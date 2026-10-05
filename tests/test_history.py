"""Venue history: address and owner matching, classification, each source's
check, score and Attio/Slack/sheet effects. No network: every source talks to
a fake HTTP client. All names and addresses are synthetic."""

from __future__ import annotations

import csv
import io
import json
import logging
import zipfile
from datetime import date
from urllib.parse import unquote

from openpyxl import load_workbook

from licmon import attio, history, leadsheet, qualify, slack
from licmon.history import (ADDING_PERMIT, NEW_OWNER, NEW_VENUE, UNKNOWN, History,
                            Prior, classify, normalize_owner, same_owner, same_premises)
from licmon.models import Record, Snapshot
from licmon.sources.ca_abc import CaAbcSource
from licmon.sources.chicago_bacp import ChicagoBacpSource
from licmon.sources.chicago_pending import ChicagoPendingSource, account_of
from licmon.sources.fl_abt import FlAbtSource
from licmon.sources.ny_sla import NySlaSource
from licmon.sources.tx_tabc import TxTabcSource
from licmon.sources.wa_lcb import WaLcbSource

TODAY = date(2026, 9, 30)


class FakeHttp:
    """Serves Socrata JSON by dataset id; records every query."""

    def __init__(self, datasets: dict[str, list[dict]]):
        self.datasets = datasets
        self.calls: list[tuple[str, dict]] = []

    def get(self, url, params=None, headers=None):
        self.calls.append((url, dict(params or {})))
        dataset = url.rsplit("/", 1)[-1].removesuffix(".json")
        rows = self.datasets.get(dataset, [])
        where = (params or {}).get("$where", "")
        if "license_id in(" in where:
            ids = where.split("license_id in(", 1)[1].split(")", 1)[0].split(",")
            rows = [r for r in rows if r["license_id"].split(".")[0] in ids]
        offset = int((params or {}).get("$offset", 0))
        return Snapshot(url=url, body=json.dumps(rows[offset:] if offset == 0 else []).encode())


def rec(source="tx_tabc_pending", srid="1", **kw) -> Record:
    base = dict(source=source, source_record_id=srid, source_url="https://example.invalid",
                legal_name="Zebra Fake Rocks LLC", dba="Zebra Fake Ballroom",
                application_type="ORIGINAL", application_date=date(2026, 9, 1),
                address="216 Fakeinth St", city="Dallas", state="TX", zip="75207",
                county="Dallas", category="on_premise")
    base.update(kw)
    return Record(**base)


# ---------------------------------------------------------------------------
# Owner names and addresses
# ---------------------------------------------------------------------------

def test_owner_names_normalize_entity_words_and_punctuation():
    assert normalize_owner("Zebra Fake Beverages, L.L.C.") == "ZEBRA FAKE BEVERAGES"
    assert normalize_owner("THE ZEBRA FAKE CO.") == "ZEBRA FAKE"
    assert normalize_owner("Fake & Sons, Inc") == normalize_owner("FAKE AND SONS INCORPORATED")
    assert normalize_owner("O'Fake's Pub Corp") == "OFAKES PUB"
    assert normalize_owner("LLC") is None and normalize_owner(None) is None
    assert same_owner("Zebra Fake LLC", None, "ZEBRA FAKE, INC.", None)
    assert not same_owner("Zebra Fake Rocks LLC", None, "Zebra Fake Beverages LLC", None)
    # the source's own owner id wins over differently written names
    assert same_owner("Zebra Holdings", "2100038270", "ZF Group", "2100038270")
    assert not same_owner(None, None, None, None)


def test_address_match_near_misses():
    # "St" vs "Street", case and punctuation
    assert same_premises("216 Fakeinth St", "75207", None, "216 FAKEINTH STREET.", "75207-1234", None)
    # a missing street type on one side still matches
    assert same_premises("216 Fakeinth", "75207", None, "216 Fakeinth St", "75207", None)
    # different house number, or one that only starts the same
    assert not same_premises("216 Fakeinth St", "75207", None, "2160 Fakeinth St", "75207", None)
    # different ZIP
    assert not same_premises("216 Fakeinth St", "75207", None, "216 Fakeinth St", "75208", None)
    # different street type at the same number is a different street
    assert not same_premises("100 Fake St", "11101", None, "100 Fake Ave", "11101", None)
    # suites must agree: a shopping-center tenant is not the anchor store
    assert same_premises("9 Fake Blvd Suite 100", "33130", None, "9 Fake Blvd #100", "33130", None)
    assert not same_premises("9 Fake Blvd Ste 100", "33130", None, "9 Fake Blvd Ste 200",
                             "33130", None)
    assert not same_premises("4301 W Fake Dr Suite B108", "78745", None, "4301 W FAKE DR",
                             "78745", None)
    # floors are ignored however they are written
    assert same_premises("5 Fake St Fl 2 Ste 300", "10001", None, "5 FAKE ST STE 300",
                         "10001", None)
    assert same_premises("5 Fake St 2nd Floor", "10001", None, "5 FAKE STREET", "10001", None)
    # directions abbreviate; a different direction is a different address
    assert same_premises("50 North Fake Ave", "60614", None, "50 N FAKE AVE", "60614", None)
    assert not same_premises("50 N Fake Ave", "60614", None, "50 S Fake Ave", "60614", None)
    # Chicago list shape ("Floor: 1", no ZIP) vs the dataset shape ("... AVE 1ST")
    assert same_premises("4500 N Fakebelieve Ave, Floor: 1", None, "CHICAGO",
                         "4500 N FAKEBELIEVE AVE 1ST", "60640", "Chicago")
    assert not same_premises("4500 N Fakebelieve Ave, Suite/Apt: 101", None, "CHICAGO",
                             "4500 N FAKEBELIEVE AVE 102", "60640", "CHICAGO")
    # Queens-style hyphenated numbers are one token
    assert same_premises("31-01 30th Ave", "11102", None, "31-01 30TH AVENUE", "11102", None)
    assert not same_premises("31-01 30th Ave", "11102", None, "31-15 30th Ave", "11102", None)
    # nothing to match on
    assert not same_premises("Pier 9", "94111", None, "Pier 9", "94111", None)
    assert not same_premises("216 Fakeinth St", None, None, "216 Fakeinth St", None, None)


# ---------------------------------------------------------------------------
# Classification
# ---------------------------------------------------------------------------

def test_classify_rules():
    filed = date(2026, 9, 1)
    other_active = Prior("L1", "Zebra Fake Beverages LLC", active=True, issued=date(2023, 3, 8))
    other_old = Prior("L0", "Old Fake Co", ended=date(2021, 9, 6), issued=date(2017, 9, 7))
    other_recent = Prior("L2", "Old Fake Co", ended=date(2025, 6, 1), issued=date(2019, 1, 1))
    mine = Prior("L3", "Zebra Fake Rocks, L.L.C.", active=True, issued=date(2024, 1, 1))
    mine_new = Prior("L4", "Zebra Fake Rocks LLC", active=True, issued=date(2026, 9, 10))
    owner = "Zebra Fake Rocks LLC"

    assert classify(owner, None, filed, [], TODAY) == History(NEW_VENUE, 0, None)
    # prior licenses that all ended more than two years ago
    assert classify(owner, None, filed, [other_old], TODAY).label == NEW_VENUE
    # current or recent under another owner
    h = classify(owner, None, filed, [other_active, other_old], TODAY)
    assert h == History(NEW_OWNER, 2, date(2017, 9, 7))
    assert classify(owner, None, filed, [other_recent], TODAY).label == NEW_OWNER
    # the same owner already licensed there: adding a permit, even beside another business
    assert classify(owner, None, filed, [mine], TODAY).label == ADDING_PERMIT
    assert classify(owner, None, filed, [mine, other_active], TODAY).label == ADDING_PERMIT
    # the venue's own license issued after this filing is not a prior one
    assert classify(owner, None, filed, [mine_new], TODAY) == History(NEW_VENUE, 0, None)
    # owner id beats names
    assert classify("Other Name", "77", filed,
                    [Prior("L5", "Different", owner_id="77", active=True)],
                    TODAY).label == ADDING_PERMIT


def test_best_label_order():
    assert history.best([ADDING_PERMIT, NEW_OWNER]) == NEW_OWNER
    assert history.best([NEW_OWNER, UNKNOWN]) == UNKNOWN
    assert history.best([UNKNOWN, NEW_VENUE, ADDING_PERMIT]) == NEW_VENUE
    assert history.best([None, "junk"]) is None


def test_lookup_failure_is_unknown_and_logs_counts_only(caplog):
    class Broken(TxTabcSource):
        def venue_history(self, http, records, snapshots=None, today=None):
            raise RuntimeError("Zebra Fake Rocks LLC at 216 Fakeinth St")

    recs = [rec(srid="1"), rec(srid="2")]
    with caplog.at_level(logging.INFO, logger="licmon"):
        out = history.lookup(Broken(), FakeHttp({}), recs, today=TODAY)
    assert {k: v.label for k, v in out.items()} == {"1": UNKNOWN, "2": UNKNOWN}
    text = caplog.text
    assert "history FAILED tx_tabc_pending (RuntimeError)" in text
    assert "Zebra" not in text and "Fakeinth" not in text
    assert "checked 2; new venue 0, new owner 0, adding a permit 0, unknown 2" in text


# ---------------------------------------------------------------------------
# Sources
# ---------------------------------------------------------------------------

def tx_license(license_id, owner, status, issued, expires, mf, address="216 Fakeinth Street",
               zip_code="75207"):
    return {"license_id": f"{license_id}.0", "primary_status": status, "owner": owner,
            "original_issue_date": f"{issued}T00:00:00.000",
            "expiration_date": f"{expires}T00:00:00.000", "master_file_id": f"{mf}.0",
            "address": address, "zip": zip_code, "city": "Dallas"}


def test_tx_new_company_at_a_licensed_address_is_new_owner():
    # Mirrors the Longhorn Ballroom case with fake names: an active MB since
    # 2023 under another company, and an expired one from 2017 to 2021.
    http = FakeHttp({"7hf9-qc9f": [
        tx_license(200087031, "Zebra Fake Beverages LLC", "Active", "2023-03-08",
                   "2027-03-07", 2100038270),
        tx_license(105415115, "Older Fake Co", "Expired - Original Required", "2017-09-07",
                   "2021-09-06", 2100020050),
        tx_license(999, "Next Door Fake LLC", "Active", "2020-01-01", "2027-01-01", 5,
                   address="2160 Fakeinth St"),
    ]})
    r = rec(raw={"primary_license_id": "200209330.0"})
    out = TxTabcSource().venue_history(http, [r], today=TODAY)
    assert out["1"] == History(NEW_OWNER, 2, date(2017, 9, 7))
    # one batched address query, filtered by ZIP and house number
    [(url, params)] = http.calls
    assert url.endswith("/7hf9-qc9f.json")
    assert unquote(params["$where"]) == ("((starts_with(zip, '75207') AND "
                                         "starts_with(address, '216')))")


def test_tx_subordinate_permit_on_an_existing_license_is_adding_a_permit():
    http = FakeHttp({"7hf9-qc9f": [
        tx_license(300, "Zebra Fake Rocks LLC", "Active", "2024-05-01", "2026-12-01", 42,
                   address="1 Other Fake Rd", zip_code="75001"),
        tx_license(301, "Zebra Fake Rocks LLC", "Active", "2026-09-15", "2028-09-15", 43,
                   address="1 Other Fake Rd", zip_code="75001"),
    ]})
    late_hours = rec(srid="2", license_type="LH", address=None, zip=None,
                     raw={"primary_license_id": "300.0", "subordinate_license_id": "310.0",
                          "master_file_id": "42.0"})
    # a new bar's own permit, issued after its late-hours filing: still new
    new_bar = rec(srid="3", license_type="FB", address=None, zip=None,
                  raw={"primary_license_id": "301.0", "subordinate_license_id": "311.0",
                       "master_file_id": "43.0"})
    no_address = rec(srid="4", address=None, zip=None, raw={})
    out = TxTabcSource().venue_history(http, [late_hours, new_bar, no_address], today=TODAY)
    assert out["2"].label == ADDING_PERMIT
    assert out["3"].label == NEW_VENUE
    assert "4" not in out  # nothing to match on: Unknown


def ny_active(pid, owner, address, zip_code, issued):
    return {"licensepermitid": pid, "legalname": owner, "actualaddressofpremises": address,
            "city": "BROOKLYN", "zipcode": zip_code, "originalissuedate": f"{issued}T00:00:00.000",
            "expirationdate": "2028-01-31T00:00:00.000"}


def test_ny_checks_active_and_inactive_lists():
    http = FakeHttp({
        "9s3h-dpkz": [ny_active("A1", "ZEBRA FAKE BAR INC", "315 FAKE AVENUE", "11208",
                                "2019-04-01")],
        "6dg3-2z7i": [{"license_permit_id": "I1", "legalname": "OLD FAKE LLC",
                       "actual_address_of_premises": "9 FAKE PL", "zip_code": "11215",
                       "original_issue_date": "2010-01-01T00:00:00.000",
                       "expiration_date": "2020-01-31T00:00:00.000"},
                      {"license_permit_id": "I2", "legalname": "LAPSED FAKE LLC",
                       "actual_address_of_premises": "12 FAKE PL", "zip_code": "11215",
                       "original_issue_date": "2015-01-01T00:00:00.000",
                       "expiration_date": "2025-12-31T00:00:00.000"}],
    })
    src = NySlaSource()
    recs = [rec("ny_sla_pending", "a", legal_name="Zebra Fake Bar, Inc.", address="315 Fake Ave",
                zip="11208", state="NY"),
            rec("ny_sla_pending", "b", legal_name="New Fake LLC", address="9 Fake Pl",
                zip="11215", state="NY"),
            rec("ny_sla_pending", "c", legal_name="New Fake LLC", address="12 Fake Place",
                zip="11215", state="NY"),
            rec("ny_sla_pending", "d", address="Fake Pier", zip="11215", state="NY")]
    out = src.venue_history(http, recs, today=TODAY)
    assert out["a"].label == ADDING_PERMIT
    assert out["b"] == History(NEW_VENUE, 1, date(2010, 1, 1))  # ended over two years ago
    assert out["c"].label == NEW_OWNER  # lapsed last year under another company
    assert "d" not in out
    assert [u.rsplit("/", 1)[-1] for u, _ in http.calls] == ["9s3h-dpkz.json", "6dg3-2z7i.json"]


def chi_row(number, account, owner, address, expires, status="AAI", start="2020-01-01"):
    return {"license_number": number, "account_number": account, "legal_name": owner,
            "address": address, "city": "CHICAGO", "zip_code": "60640",
            "license_status": status, "license_start_date": f"{start}T00:00:00.000",
            "expiration_date": f"{expires}T00:00:00.000"}


def test_chicago_filing_types_and_address_check():
    http = FakeHttp({"r5kz-chrr": [
        # the same license over two terms counts once
        chi_row("111", "9001", "OLD FAKE TAVERN INC", "4500 N FAKEBELIEVE AVE 1ST", "2025-02-15",
                start="2021-02-16"),
        chi_row("111", "9001", "OLD FAKE TAVERN INC", "4500 N FAKEBELIEVE AVE 1ST", "2027-02-15",
                start="2025-02-16"),
        chi_row("222", "9002", "ZEBRA FAKE LLC", "10 W FAKE ST", "2027-05-15"),
    ]})
    src = ChicagoBacpSource()
    expansion = rec("chicago_bacp_liquor", "e", application_type="EXPANSION",
                    address="10 W FAKE ST", zip="60640", state="IL", city="CHICAGO")
    issue = rec("chicago_bacp_liquor", "i", legal_name="New Fake LLC", application_type="NEW",
                address="4500 N FAKEBELIEVE AVE", zip="60640", state="IL", city="CHICAGO",
                raw={"account_number": "9100", "license_number": "333"})
    own = rec("chicago_bacp_liquor", "o", legal_name="Zebra Fake LLC", application_type="NEW",
              address="10 W FAKE ST", zip="60640", state="IL", city="CHICAGO",
              raw={"account_number": "9002", "license_number": "222"})
    out = src.venue_history(http, [expansion, issue, own], today=TODAY)
    assert out["e"].label == ADDING_PERMIT  # straight from the filing type
    assert out["i"] == History(NEW_OWNER, 1, date(2021, 2, 16))
    assert out["o"] == History(NEW_VENUE, 0, None)  # its own license is not a prior one


def test_chicago_pending_matches_by_city_and_account():
    assert account_of("12345-2-abcdef123456") == "12345"
    assert account_of("noid-abcdef123456") is None
    http = FakeHttp({"r5kz-chrr": [
        chi_row("111", "12345", "SOME OTHER NAME LLC", "4500 N FAKEBELIEVE AVE 1ST",
                "2027-02-15")]})
    pending = rec("chicago_bacp_pending", "12345-2-abcdef123456", legal_name="Zebra Fake LLC",
                  application_type="NEW", address="4500 N Fakebelieve Ave, Floor: 1",
                  zip=None, state="IL", city="CHICAGO")
    out = ChicagoPendingSource().venue_history(http, [pending], today=TODAY)
    assert out[pending.source_record_id].label == ADDING_PERMIT  # same BACP account
    [(_, params)] = http.calls
    assert "upper(city) = 'CHICAGO'" in params["$where"]
    assert "license_code in (" in params["$where"]


FL_HEADER = ["Board", "Profession", "Owner Name", "Series", "Modifier", "Mail Address 1",
             "Mail Address 2", "Mail Address 3", "Mail City", "Mail State", "Mail ZIP",
             "Mail County", "DBA", "Location Address 1", "Location Address 2",
             "Location Address 3", "Location City", "Location State", "Location ZIP",
             "Location County", "License Number", "Primary Status", "Secondary Status",
             "Original Licensure Date", "Effective Date", "Expiration Date",
             "Tax Stamp Designation", "Smoking Designation", "Retail Tobacco Indicator"]


def fl_row(number, owner, address, primary, secondary, orig, expires, zip_code="33130"):
    row = dict.fromkeys(FL_HEADER, "")
    row.update({"Owner Name": owner, "Series": "4COP", "Location Address 1": address,
                "Location City": "MIAMI", "Location State": "FL", "Location ZIP": zip_code,
                "License Number": number, "Primary Status": primary,
                "Secondary Status": secondary, "Original Licensure Date": orig,
                "Expiration Date": expires})
    return row


def fl_snapshot(rows) -> Snapshot:
    buf = io.StringIO()
    w = csv.DictWriter(buf, FL_HEADER, quoting=csv.QUOTE_ALL)
    w.writeheader()
    w.writerows(rows)
    return Snapshot(url="https://example.invalid/fl.csv", body=buf.getvalue().encode("latin-1"))


def test_fl_uses_its_own_extract_including_transfers():
    snap = fl_snapshot([
        # transfer in progress: outgoing licensee, then the incoming one
        fl_row("BEV1", "OLD FAKE LOUNGE LLC", "900 FAKE BAY TER", "20", "35", "01/05/2015",
               "03/31/2027"),
        fl_row("BEV1", "ZEBRA FAKE LOUNGE LLC", "900 FAKE BAY TER", "21", "20", "09/01/2026",
               "03/31/2027"),
        # the same owner already holds another license in the same suite
        fl_row("BEV2", "ZEBRA FAKE GROUP INC", "12 FAKE AVE STE 3", "20", "20", "02/02/2020",
               "03/31/2027"),
        fl_row("BEV3", "ZEBRA FAKE GROUP INC", "12 FAKE AVENUE #3", "20", "20", "09/10/2026",
               "03/31/2027"),
        # a new license at a fresh address
        fl_row("BEV4", "BRAND NEW FAKE LLC", "77 FAKE CT", "20", "20", "09/12/2026",
               "03/31/2027"),
    ])
    src = FlAbtSource()
    recs = {r.source_record_id: r for r in src.parse([snap])}
    out = src.venue_history(None, list(recs.values()), snapshots=[snap], today=TODAY)
    assert out["BEV1"] == History(NEW_OWNER, 1, date(2015, 1, 5))
    assert out["BEV3"].label == ADDING_PERMIT
    assert out["BEV4"] == History(NEW_VENUE, 0, None)
    # without this run's download (requalify --history) it fetches once
    class Http:
        calls = 0

        def get(self, url, params=None, headers=None):
            Http.calls += 1
            return snap

    src.venue_history(Http(), [recs["BEV4"]], snapshots=None, today=TODAY)
    assert Http.calls == 1


CA_HEADER = ["License Type", "File Number", "Lic or App", "Type Status", "Type Orig Iss Date",
             "Expir Date", "Fee Codes", "Dup Counts", "Master Ind", "Term in # of Months",
             "Geo Code", "District", "Primary Name", "Prem Addr 1", "Prem Addr 2", "Prem City",
             "Prem State", "Prem Zip", "DBA Name", "Mail Addr 1", "Mail Addr 2", "Mail City",
             "Mail State", "Mail Zip", "Prem County", "Prem Census Tract #"]


def ca_row(file_number, kind, status, owner, address, orig="", expires="", zip_code="90012-1234"):
    row = dict.fromkeys(CA_HEADER, " ")
    row.update({"License Type": "47", "File Number": file_number, "Lic or App": kind,
                "Type Status": status, "Type Orig Iss Date": orig, "Expir Date": expires,
                "Primary Name": owner, "Prem Addr 1": address, "Prem City": "LOS ANGELES",
                "Prem State": "CA", "Prem Zip": zip_code, "Prem County": "LOS ANGELES"})
    return row


def ca_snapshot(rows) -> Snapshot:
    buf = io.StringIO()
    buf.write("﻿\"Updated Tuesday 30th of September 2026 03:50:28 AM\"\r\n")
    w = csv.DictWriter(buf, CA_HEADER, quoting=csv.QUOTE_ALL)
    w.writeheader()
    w.writerows(rows)
    zbuf = io.BytesIO()
    with zipfile.ZipFile(zbuf, "w") as z:
        z.writestr("ABC-DailyDataExport.csv", buf.getvalue().encode("utf-8"))
    return Snapshot(url="https://example.invalid/ca.zip", body=zbuf.getvalue())


def test_ca_compares_applications_with_license_rows():
    snap = ca_snapshot([
        ca_row("0600001", "APP", "PEND", "ZEBRA FAKE LLC", "500 S FAKE ST"),
        ca_row("0500001", "LIC", "ACTIVE", "OLD FAKE INC", "500 SOUTH FAKE STREET",
               "12-JUN-2018", "30-JUN-2027"),
        # its own file number, same owner: the application's own license
        ca_row("0600002", "APP", "ACTIVE", "SAME FAKE LLC", "7 FAKE WAY"),
        ca_row("0600002", "LIC", "ACTIVE", "SAME FAKE LLC", "7 FAKE WAY", "01-SEP-2026",
               "30-SEP-2027"),
        # surrendered long ago
        ca_row("0600003", "APP", "PEND", "NEWER FAKE LLC", "8 FAKE WAY"),
        ca_row("0400009", "LIC", "SUREND", "GONE FAKE LLC", "8 FAKE WAY", "01-MAR-2001",
               "31-MAR-2019"),
    ])
    src = CaAbcSource()
    recs = list(src.parse([snap]))
    assert len(recs) == 3  # LIC rows are still not leads
    out = src.venue_history(None, recs, snapshots=[snap], today=TODAY)
    assert out["0600001"] == History(NEW_OWNER, 1, date(2018, 6, 12))
    assert out["0600002"] == History(NEW_VENUE, 0, None)
    assert out["0600003"] == History(NEW_VENUE, 1, date(2001, 3, 1))


def test_wa_uses_its_own_application_type():
    src = WaLcbSource()
    recs = [rec("wa_lcb_actions", "a", application_type="ASSUMPTION"),
            rec("wa_lcb_actions", "b", application_type="ADDED/CHANGE OF CLASS/IN LIEU"),
            rec("wa_lcb_actions", "c", application_type="NEW APPLICATION"),
            rec("wa_lcb_actions", "d", application_type="CHANGE OF LOCATION")]
    out = src.venue_history(None, recs)
    assert {k: v.label for k, v in out.items()} == {"a": NEW_OWNER, "b": ADDING_PERMIT}


# ---------------------------------------------------------------------------
# Score, sheet, Attio, Slack
# ---------------------------------------------------------------------------

def test_score_adjustments():
    lounge = rec(dba="Zebra Fake Lounge", legal_name="Zebra Fake LLC", category="nightlife",
                 license_type="LH", license_description="Late Hours Certificate",
                 status="Pending - In Review")
    base = qualify.qualify(lounge, "Dallas", today=TODAY)
    assert base.qualified and base.tier == "A" and base.hot
    owner = qualify.qualify(lounge, "Dallas", today=TODAY, history=NEW_OWNER)
    assert owner.lead_score == base.lead_score - 10 and not owner.hot
    permit = qualify.qualify(lounge, "Dallas", today=TODAY, history=ADDING_PERMIT)
    assert permit.lead_score == base.lead_score - 15 and not permit.hot
    for label in (NEW_VENUE, UNKNOWN, None):
        same = qualify.qualify(lounge, "Dallas", today=TODAY, history=label)
        assert (same.lead_score, same.hot) == (base.lead_score, base.hot)
    # New owner can never be Hot, even with a threshold it would clear
    assert not qualify.apply_history(qualify.Qualification(True, 99, "A", "", lead_score=100,
                                                           hot=True), NEW_OWNER).hot


def sheet_rec(record_id, label, **kw):
    base = {"queue_date": TODAY, "record_id": record_id, "venue_key": f"k{record_id}",
            "tier": "A", "score": 90, "legal_name": None, "dba": f"ZEBRA FAKE LOUNGE {record_id}",
            "license_description": "Mixed Beverage Permit", "application_type": "ORIGINAL",
            "status": "Received", "application_date": date(2026, 9, 20),
            "address": "1 FAKE ST", "city": "AUSTIN", "state": "TX", "zip": "78701",
            "metro": "Austin", "source": "tx_tabc_pending",
            "source_url": "https://example.invalid/r", "category": "on_premise", "raw": {},
            "lead_score": 70, "venue_history": label}
    base.update(kw)
    return base


def test_existing_venues_tab_and_column():
    rows = leadsheet.group_records([
        sheet_rec(1, NEW_VENUE), sheet_rec(2, None), sheet_rec(3, NEW_OWNER),
        sheet_rec(4, ADDING_PERMIT),
        # one venue, two filings: the best label wins
        sheet_rec(5, ADDING_PERMIT, venue_key="k5"), sheet_rec(6, NEW_OWNER, venue_key="k5")])
    labels = {r["lead_ids"]: r["venue_history"] for r in rows}
    assert labels == {"1": NEW_VENUE, "2": UNKNOWN, "3": NEW_OWNER, "4": ADDING_PERMIT,
                      "5 6": NEW_OWNER}
    wb = load_workbook(io.BytesIO(leadsheet.build_workbook(rows, rows)))
    assert wb.sheetnames[:3] == ["New", "Existing venues", "All open"]
    col = leadsheet.HEADERS.index("Venue history") + 1
    new = wb["New"]
    assert [new.cell(row=r, column=col).value for r in (2, 3)] == [NEW_VENUE, UNKNOWN]
    assert new.max_row == 3
    existing = wb["Existing venues"]
    assert existing.cell(row=1, column=1).value == leadsheet.EXISTING_NOTE
    assert [c.value for c in existing[2]] == leadsheet.HEADERS
    assert sorted(existing.cell(row=r, column=col).value for r in (3, 4, 5)) == \
        [ADDING_PERMIT, NEW_OWNER, NEW_OWNER]
    assert existing.max_row == 5 and existing.freeze_panes == "E3"
    # open and state tabs keep everything, with the column
    open_headers = [c.value for c in wb["All open"][1]]
    assert "Venue history" in open_headers and wb["All open"].max_row == 6
    assert wb["TX"].max_row == 6
    legend = " ".join(str(v) for row in wb["How scoring works"].values for v in row if v)
    assert "Existing venues" in legend and "Adding a permit" in legend
    assert "-10" in legend and "-15" in legend
    assert "—" not in legend + leadsheet.EXISTING_NOTE


def attio_row(i, label, **kw):
    base = {"venue_key": f"TX|78701|{i} FAKE ST", "priority": "A", "hot": "", "lead_score": 70,
            "business_name": f"Zebra Fake Lounge {i}", "stage": "In review", "state": "TX",
            "venue_history": label}
    base.update(kw)
    return base


class FakeResp:
    def __init__(self, status=200, body=None):
        self.status_code = status
        self._body = body if body is not None else {"data": {}}
        self.headers = {}

    def json(self):
        return self._body


class AttioSession:
    """Minimal fake Attio for the Venue history field."""

    def __init__(self, fields=("venue_key", "priority"), refuse_fields=False):
        self.headers = {}
        self.calls = []
        self.fields = list(fields)
        self.refuse_fields = refuse_fields

    def request(self, method, url, json=None, params=None, timeout=None):
        path = url.removeprefix(attio.API)
        self.calls.append((method, path, json))
        if path == "/lists/license_leads/entries/query":
            return FakeResp(body={"data": []})
        if path == "/objects/target_client/records/query":
            return FakeResp(body={"data": [{"id": {"record_id": "rec-team"},
                                            "values": {"company_1": [
                                                {"value": json["filter"]["company_1"]["value"]
                                                 ["$contains"]}]}}]})
        if path == "/lists/license_leads/attributes/priority/options":
            return FakeResp(body={"data": [{"title": t} for t in ("Hot", "A", "B")]})
        if path == "/lists/license_leads/attributes":
            if method == "POST":
                if self.refuse_fields:
                    return FakeResp(403, {"code": "missing_scope"})
                self.fields.append(json["data"]["api_slug"])
                return FakeResp(body={"data": {}})
            return FakeResp(body={"data": [{"api_slug": f} for f in self.fields]})
        return FakeResp()

    def posts(self, path):
        return [body for m, p, body in self.calls if m == "POST" and p == path]


def attio_client(session):
    return attio.Client(key="test-key", session=session, sleep=lambda s: None)


def test_attio_skips_permits_labels_new_owner_and_adds_the_field():
    rows = [attio_row(1, NEW_OWNER), attio_row(2, ADDING_PERMIT, hot="Hot", lead_score=95),
            attio_row(3, NEW_VENUE)]
    assert [r["venue_key"] for r in attio.candidates(rows, 60)] == [
        "TX|78701|1 FAKE ST", "TX|78701|3 FAKE ST"]
    s = AttioSession()
    counts = attio.sync(attio_client(s), rows, write=True, cap=50, min_b=60)
    assert counts["permits_left_out"] == 1 and counts["added"] == 2
    assert counts["history_field_added"] == 1 and not counts["history_field_missing"]
    [field] = s.posts("/lists/license_leads/attributes")
    assert field["data"]["api_slug"] == "venue_history" and field["data"]["type"] == "text"
    # the field is made before the first entry is written
    first_entry = next(i for i, c in enumerate(s.calls)
                       if c[0] == "POST" and c[1] == "/lists/license_leads/entries")
    first_field = next(i for i, c in enumerate(s.calls)
                       if c[0] == "POST" and c[1] == "/lists/license_leads/attributes")
    assert first_field < first_entry
    values = {e["data"]["entry_values"]["venue_key"]: e["data"]["entry_values"]
              for e in s.posts("/lists/license_leads/entries")}
    assert values["TX|78701|1 FAKE ST"]["venue_history"] == NEW_OWNER
    assert values["TX|78701|3 FAKE ST"]["venue_history"] == NEW_VENUE
    assert attio.update_values(rows[0])["venue_history"] == NEW_OWNER
    # next run: the field exists, nothing added
    s2 = AttioSession(fields=("venue_key", "priority", "venue_history"))
    counts = attio.sync(attio_client(s2), rows, write=True, cap=50, min_b=60)
    assert counts["history_field_added"] == 0 and s2.posts("/lists/license_leads/attributes") == []


def test_attio_without_schema_scope_syncs_without_the_field():
    s = AttioSession(refuse_fields=True)
    counts = attio.sync(attio_client(s), [attio_row(1, NEW_OWNER)], write=True, cap=50)
    assert counts["history_field_missing"] and counts["added"] == 1
    [entry] = s.posts("/lists/license_leads/entries")
    assert "venue_history" not in entry["data"]["entry_values"]


def test_attio_dry_run_never_adds_the_field():
    s = AttioSession()
    counts = attio.sync(attio_client(s), [attio_row(1, NEW_OWNER)], write=False, cap=50)
    assert counts["history_field_added"] == 1  # would add it
    assert [c for c in s.calls if c[0] == "POST" and c[1] != "/lists/license_leads/entries/query"
            and not c[1].endswith("/records/query")] == []


def test_slack_never_names_permits_and_labels_new_owner():
    rows = [attio_row(1, ADDING_PERMIT, hot="Hot", lead_score=90, city="Austin",
                      whats_new=leadsheet.NEW_FILING),
            attio_row(2, NEW_OWNER, hot="Hot", lead_score=80, city="Austin",
                      whats_new=leadsheet.NEW_FILING),
            attio_row(3, NEW_VENUE, hot="Hot", lead_score=85, city="Austin")]
    text = slack.compose(rows)
    assert "Zebra Fake Lounge 1" not in text
    assert "Zebra Fake Lounge 2, Austin (In review, New owner)" in text
    assert "Zebra Fake Lounge 3, Austin (In review)" in text


def test_same_premises_suite_one_side():
    # Venue history keeps the strict rule; the contact lookup relaxes it.
    assert not same_premises("4301 Fake Dr", "75001", None, "4301 Fake Dr Ste B108", "75001", None)
    assert same_premises("4301 Fake Dr", "75001", None, "4301 Fake Dr Ste B108", "75001", None,
                         suite_one_side_ok=True)
    assert not same_premises("4301 Fake Dr Ste A", "75001", None, "4301 Fake Dr Ste B108", "75001", None,
                             suite_one_side_ok=True)
