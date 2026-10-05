"""The owner's lead spreadsheet. Synthetic data only."""

import csv
import io
from datetime import date

from openpyxl import load_workbook

from licmon import leadsheet
from licmon.sources.ca_abc import CaAbcSource
from licmon.sources.fl_abt import FlAbtSource
from licmon.sources.wa_lcb import WaLcbSource

DAY = date(2026, 10, 1)


def rec(record_id, **kw):
    base = {
        "queue_date": DAY, "record_id": record_id, "venue_key": f"k{record_id}",
        "tier": "B", "score": 60, "legal_name": None, "dba": None,
        "license_description": "Mixed Beverage Permit", "application_type": "ORIGINAL",
        "status": "Received", "application_date": date(2026, 9, 20),
        "address": "1 FAKE ST", "city": "AUSTIN", "state": "TX", "zip": "78701",
        "metro": "Austin", "source": "tx_tabc_pending",
        "source_url": "https://example.invalid/r", "category": "on_premise", "raw": {},
    }
    base.update(kw)
    return base


def test_contact_hooks_use_only_the_official_row():
    wa = WaLcbSource().contact({"Contact Phone": "2065550100",
                                "Applicant(s)": "ZEBRA FAKE LLC; JANE Q TESTER"})
    assert wa == {"phone": "2065550100", "people": "ZEBRA FAKE LLC; JANE Q TESTER"}
    ca = CaAbcSource().contact({"rows": [
        {"Mail Addr 1": "", "Mail City": ""},
        {"Mail Addr 1": "PO BOX 1", "Mail Addr 2": "", "Mail City": "FAKEVILLE",
         "Mail State": "CA", "Mail Zip": "90000"}]})
    assert ca == {"mailing_address": "PO BOX 1, FAKEVILLE, CA 90000"}
    fl = FlAbtSource().contact({"Mail Address 1": "9 TEST WAY", "Mail Address 2": "STE 2",
                                "Mail City": "MIAMI", "Mail State": "FL", "Mail ZIP": "33101"})
    assert fl == {"mailing_address": "9 TEST WAY STE 2, MIAMI, FL 33101"}
    assert FlAbtSource().contact({}) == {}


def test_group_merges_venue_and_cleans_fields():
    rows = leadsheet.group_records([
        rec(1, venue_key="v", tier="A", score=80, legal_name="ZEBRA FAKE LLC",
            dba="ZEBRA FAKE LOUNGE", license_description="Late Hours Certificate",
            category="nightlife", application_date=date(2026, 9, 18)),
        rec(2, venue_key="v", tier="B", score=60, legal_name="ZEBRA FAKE LLC",
            dba="ZEBRA FAKE LOUNGE", status="Pending – In Review"),
        rec(3, legal_name="Other Fake Inc", dba=None, tier="C", score=45),
    ])
    assert [r["business_name"] for r in rows] == ["Zebra Fake Lounge", "Other Fake INC"]
    top = rows[0]
    assert top["priority"] == "A"
    assert top["company"] == "Zebra Fake LLC"
    assert top["business_type"] == "Nightclub / lounge / ticketed venue"
    assert top["filing"] == "New application"
    assert top["stage"] == "In review"
    assert top["filed_on"] == date(2026, 9, 18)
    assert top["license"] == "Late Hours Certificate; Mixed Beverage Permit"
    assert top["lead_ids"] == "1 2"
    assert top["map_url"].startswith("https://www.google.com/maps/search/?api=1&query=ZEBRA")
    assert top["google_url"] == "https://www.google.com/search?q=ZEBRA+FAKE+LOUNGE+AUSTIN+TX"
    assert top["instagram_url"] == ("https://www.google.com/search?q="
                                    "site%3Ainstagram.com+ZEBRA+FAKE+LOUNGE")
    assert rows[1]["company"] is None  # no DBA: the name is the company


def test_wa_phone_and_people_without_repeating_business():
    [row] = leadsheet.group_records([rec(
        7, source="wa_lcb_actions", dba="ZEBRA FAKE BAR", legal_name=None,
        application_type="ASSUMPTION", status="APPROVED",
        raw={"Contact Phone": "206-555-0100",
             "Applicant(s)": "ZEBRA FAKE BAR; JANE Q TESTER; JOHN TESTER"})])
    assert row["phone"] == "(206) 555-0100"
    assert row["people"] == "Jane Q Tester; John Tester"
    assert row["filing"] == "Change of owner"
    assert row["stage"] == "Licensed"  # WA "approved" section = issued


def test_filing_and_status_labels():
    assert leadsheet._filing(["NEW LICENSE"], "fl_abt_licenses") == "Newly licensed"
    assert leadsheet._filing([], "ca_abc_applications") == "New application"
    assert leadsheet._filing(["ADDED/CHANGE OF TRADENAME"], "x") == "Name change"
    by_source = {"fl": FlAbtSource(), "ca": CaAbcSource()}

    def stage_of(status, source):
        return leadsheet._stage([{"status": status, "source": source}], by_source)

    assert stage_of("Current", "fl") == "Licensed"
    assert stage_of("PEND", "ca") == "Received"
    assert stage_of("IntakeComplete", "unknown") == "Received"
    assert stage_of("Withdrawn (application withdrawn)", "fl") == ""
    # a stored stage wins over the status text
    assert leadsheet._stage([{"status": "PEND", "stage": "Approved", "source": "ca"}],
                            by_source) == "Approved"


def test_xlsx_layout_links_and_no_formulas():
    rows = leadsheet.group_records([
        rec(1, tier="A", score=80, dba="=HYPERLINK(\"http://evil.invalid\")"),
        rec(2, dba="ZEBRA FAKE TAVERN")])
    wb = load_workbook(io.BytesIO(leadsheet.build_xlsx(rows, title="Leads 2026-10-01")))
    ws = wb.active
    assert ws.title == "Leads 2026-10-01"
    assert [c.value for c in ws[1]] == leadsheet.HEADERS
    assert ws.freeze_panes == "E2"  # through Business name
    assert ws.auto_filter.ref == "A1:AC3"  # 25 columns plus Contact person and 3 links
    name_col = leadsheet.HEADERS.index("Business name") + 1
    evil = ws.cell(row=2, column=name_col)
    assert evil.data_type == "s" and evil.value.startswith("=")
    rec_col = leadsheet.HEADERS.index("Official record") + 1
    assert ws.cell(row=2, column=rec_col).value == "Record"
    assert ws.cell(row=2, column=rec_col).hyperlink.target == "https://example.invalid/r"
    date_col = leadsheet.HEADERS.index("Filed on") + 1
    assert ws.cell(row=2, column=date_col).value.date() == date(2026, 9, 20)


def test_csv_has_headers_and_neutralizes_formulas():
    rows = leadsheet.group_records([rec(1, dba="=1+1")])
    buf = io.StringIO()
    leadsheet.write_csv(rows, buf)
    out = list(csv.reader(io.StringIO(buf.getvalue())))
    assert out[0] == leadsheet.HEADERS
    assert out[1][leadsheet.HEADERS.index("Business name")] == "'=1+1"


def test_empty_sheet_is_valid():
    wb = load_workbook(io.BytesIO(leadsheet.build_xlsx([])))
    assert [c.value for c in wb.active[1]] == leadsheet.HEADERS


def test_score_hot_and_whats_new_per_venue():
    rows = leadsheet.group_records([
        rec(1, venue_key="v", tier="A", lead_score=85, hot=True, stage="Approved",
            event_type="changed", changes={"stage": ["Received", "Approved"]}),
        rec(2, venue_key="v", tier="A", lead_score=60, hot=False, stage="Received",
            event_type="changed", changes={"status": ["A", "B"]}),
        rec(3, venue_key="w", tier="B", lead_score=40, event_type="changed",
            changes={"address": ["1 A ST", "2 A ST"]}),
        rec(4, venue_key="x", tier="C", lead_score=90, event_type="new"),
    ])
    by_key = {r["venue_key"]: r for r in rows}
    assert by_key["v"]["lead_score"] == 85 and by_key["v"]["hot"] == "Hot"
    assert by_key["v"]["stage"] == "Approved"
    assert by_key["v"]["whats_new"] == leadsheet.STAGE_ADVANCED
    assert by_key["w"]["whats_new"] == leadsheet.DETAILS_CHANGED
    assert by_key["x"]["whats_new"] == leadsheet.NEW_FILING
    # sorted by score, highest first, whatever the tier
    assert [r["venue_key"] for r in rows] == ["x", "v", "w"]


def test_workbook_tabs_come_from_the_data():
    new = leadsheet.group_records([
        rec(1, tier="A", lead_score=80, hot=True, dba="ZEBRA FAKE LOUNGE", event_type="new"),
        rec(2, tier="B", lead_score=50, dba="ZEBRA FAKE TAVERN", event_type="new")])
    open_rows = leadsheet.group_records([
        rec(1, tier="A", lead_score=80, hot=True, dba="ZEBRA FAKE LOUNGE"),
        rec(5, tier="C", lead_score=30, dba="FAKE BISTRO", state="CA", city="LA",
            queue_date=date(2026, 9, 28)),
        rec(6, tier="B", lead_score=95, dba="FAKE TAPROOM", state="WA",
            queue_date=date(2026, 9, 27))])
    wb = load_workbook(io.BytesIO(leadsheet.build_workbook(new, open_rows)))
    assert wb.sheetnames == ["New", "Existing venues", "All open", "CA", "TX", "WA",
                             "How scoring works"]
    new_ws, open_ws = wb["New"], wb["All open"]
    assert [c.value for c in new_ws[1]] == leadsheet.HEADERS
    assert "What's new" in leadsheet.HEADERS and "Stage" in leadsheet.HEADERS
    open_headers = [c.value for c in open_ws[1]]
    assert "Queued on" in open_headers and "What's new" not in open_headers
    hot_col = leadsheet.HEADERS.index("Hot") + 1
    score_col = leadsheet.HEADERS.index("Score") + 1
    assert new_ws.cell(row=2, column=hot_col).value == "Hot"
    assert [new_ws.cell(row=r, column=score_col).value for r in (2, 3)] == [80, 50]
    assert new_ws.cell(row=2, column=leadsheet.HEADERS.index("What's new") + 1).value \
        == leadsheet.NEW_FILING
    # All open: every day, highest score first
    names = [open_ws.cell(row=r, column=open_headers.index("Business name") + 1).value
             for r in (2, 3, 4)]
    assert names == ["Fake Taproom", "Zebra Fake Lounge", "Fake Bistro"]
    assert wb["CA"].max_row == 2 and wb["TX"].max_row == 2
    legend = [c.value for c in wb["How scoring works"]["A"]]
    assert "Public place of amusement (Chicago)" in legend and "Licensed" in legend
    text = " ".join(str(v) for row in wb["How scoring works"].values for v in row if v)
    assert "—" not in text  # no em dashes in owner-facing text
    assert ("A: nightclubs, lounges and ticketed venues (comedy, live music, sports, "
            "theaters, event venues)") in legend
    assert "(adult) in Business type" in legend
    assert any(str(v).startswith("Theater or performing arts") for v in legend)


def test_adult_venue_is_marked_in_business_type():
    [row] = leadsheet.group_records([rec(1, tier="A", dba="ZEBRA FAKE CLUB", adult=True)])
    assert row["adult"] is True
    assert row["business_type"] == "Nightclub / lounge / ticketed venue (adult)"
    [row] = leadsheet.group_records([rec(1, tier="B", dba="ZEBRA FAKE BAR")])
    assert row["adult"] is False and row["business_type"] == "Bar / event venue"


def test_workbook_without_open_rows_is_valid():
    wb = load_workbook(io.BytesIO(leadsheet.build_workbook([])))
    assert wb.sheetnames == ["New", "Existing venues", "All open", "How scoring works"]


# --- contact lookup on the sheet (contact.py results; all invented) ---

def check(status, **kw):
    """A contact_checks row as attach_contacts reads it."""
    reachable = status == "reachable"
    base = {
        "status": status, "confidence_score": 75 if reachable else 15,
        "confidence_label": "Verified" if reachable else "Unverified",
        "confidence_reason": ("Google listing at the same address, same name; the website "
                              "links this Instagram") if reachable else
                             "Phone from the license filing (may be a lawyer)",
        "outreach_method": "Instagram DM" if reachable else "Wait: no verified contact yet",
        "outreach_second": "Call" if reachable else None,
        "contact_kind": "instagram" if reachable else "filing_phone",
        "contact_value": "@zebrafakelounge" if reachable else "(512) 555-0199",
        "contact_url": ("https://www.instagram.com/zebrafakelounge/" if reachable
                        else "tel:+15125550199"),
        "maps_url": None, "last_checked_at": date(2026, 10, 1),
        "next_check_at": None if reachable else date(2026, 10, 8),
        "newly_reachable_on": None,
        "channels": ([
            {"kind": "instagram", "value": "@zebrafakelounge",
             "url": "https://www.instagram.com/zebrafakelounge/",
             "signals": ["listing_address", "listing_name", "linked_from_website"],
             "score": 75, "label": "Verified", "last_post": "2026-09-30"},
            {"kind": "website", "value": "zebrafake.test", "url": "https://zebrafake.test/",
             "signals": ["listing_address", "listing_name", "website_loads"],
             "score": 65, "label": "Likely"},
            {"kind": "phone", "value": "(512) 555-0142", "url": "tel:+15125550142",
             "signals": ["listing_address", "listing_name", "listing_open"],
             "score": 65, "label": "Likely"},
        ] if reachable else [
            {"kind": "filing_phone", "value": "(512) 555-0199", "url": "tel:+15125550199",
             "signals": ["filing_phone"], "score": 15, "label": "Unverified"},
        ]),
    }
    base.update(kw)
    return base


def checked_rows():
    rows = leadsheet.group_records([
        rec(1, venue_key="reach", tier="A", lead_score=70, dba="ZEBRA FAKE LOUNGE",
            event_type="new"),
        rec(2, venue_key="wait", tier="A", lead_score=90, dba="QUOKKA FAKE CLUB",
            event_type="new"),
        rec(3, venue_key="never", tier="C", lead_score=30, dba="FAKE BISTRO",
            event_type="new"),
        rec(4, venue_key="newly", tier="B", lead_score=61, dba="FAKE TAPROOM",
            queue_date=date(2026, 9, 20), event_type="new"),
    ])
    by_key = {r["venue_key"]: r for r in rows}
    leadsheet.apply_contact(by_key["reach"], check("reachable"), DAY)
    leadsheet.apply_contact(by_key["wait"], check("waiting"), DAY)
    leadsheet.apply_contact(by_key["newly"], check("reachable", newly_reachable_on=DAY), DAY)
    return rows


def _table(ws, header_row=1):
    headers = [c.value for c in ws[header_row]]
    return headers, [{h: ws.cell(row=r, column=j + 1) for j, h in enumerate(headers)}
                     for r in range(header_row + 1, ws.max_row + 1)]


def test_sheet_without_contact_lookup_is_laid_out_as_before():
    rows = leadsheet.group_records([rec(1, tier="A", dba="ZEBRA FAKE LOUNGE"),
                                    rec(2, tier="B", dba="QUOKKA FAKE TAVERN")])
    wb = load_workbook(io.BytesIO(leadsheet.build_workbook(rows, rows)))
    assert wb.sheetnames == ["New", "Existing venues", "All open", "TX", "How scoring works"]
    headers, cells = _table(wb["New"])
    assert headers == leadsheet.HEADERS
    assert "Best way to reach" not in headers and "Phone" in headers
    assert len(cells) == 2  # nothing hidden when nothing was checked
    assert cells[0]["Instagram"].value == "Search"
    legend = [c.value for c in wb["How scoring works"]["A"]]
    assert "Contact confidence" not in legend


def test_sheet_moves_unreachable_venues_to_waiting_tab():
    rows = checked_rows()
    wb = load_workbook(io.BytesIO(leadsheet.build_workbook(rows, rows)))
    assert wb.sheetnames == ["Today's plan", "New", "Existing venues", "Waiting on contact",
                             "All open", "TX", "How scoring works"]
    headers, cells = _table(wb["New"])
    name = headers.index("Business name")
    assert headers[name + 1:name + 7] == ["Newly reachable", "Best way to reach", "Contact",
                                          "Confidence", "Why we trust it",
                                          "Second best way"]
    assert "Filing phone (may be a lawyer)" in headers and "Phone" not in headers
    assert headers.index("Facebook") == headers.index("Instagram") + 1
    names = [c["Business name"].value for c in cells]
    # Newly reachable first, then by score; the waiting venue is gone; the
    # venue the lookup never checked (tier C) stays.
    assert names == ["Fake Taproom", "Zebra Fake Lounge", "Fake Bistro"]
    newly, zebra, bistro = cells
    assert newly["Newly reachable"].value == "Newly reachable"
    assert zebra["Newly reachable"].value is None
    assert zebra["Best way to reach"].value == "Instagram DM"
    assert zebra["Contact"].value == "@zebrafakelounge"
    assert zebra["Contact"].hyperlink.target == "https://www.instagram.com/zebrafakelounge/"
    assert zebra["Confidence"].value == "Verified 75"
    assert zebra["Why we trust it"].value.startswith("Google listing at the same address")
    # Verified links replace the search links on the lead tabs.
    assert zebra["Instagram"].value == "Profile"
    assert zebra["Instagram"].hyperlink.target == "https://www.instagram.com/zebrafakelounge/"
    assert zebra["Google"].value == "Website"
    assert zebra["Google"].hyperlink.target == "https://zebrafake.test/"
    assert bistro["Instagram"].value == "Search" and bistro["Best way to reach"].value is None

    wait = wb["Waiting on contact"]
    assert wait.cell(row=1, column=1).value.startswith("No verified contact yet.")
    w_headers, w_cells = _table(wait, header_row=2)
    assert w_headers == [h for _, h, _ in leadsheet.WAITING_COLUMNS]
    [quokka] = w_cells
    assert quokka["Business name"].value == "Quokka Fake Club"
    assert quokka["Status"].value == "Waiting"
    assert quokka["What we found (not verified)"].value == (
        "Filing phone (may be a lawyer): (512) 555-0199 (Unverified 15)")
    assert quokka["Next check"].value.date() == date(2026, 10, 8)
    assert quokka["Instagram"].value == "Search"  # search links stay here
    assert "quokka" in quokka["Instagram"].hyperlink.target.lower()
    open_names = [c["Business name"].value for c in _table(wb["All open"])[1]]
    assert "Quokka Fake Club" not in open_names
    legend = [c.value for c in wb["How scoring works"]["A"]]
    assert "Contact confidence" in legend and "Waiting on contact" in legend
    assert "1. Instagram DM" in legend and "3. Facebook message" in legend
    text = " ".join(str(v) for ws in wb for row in ws.values for v in row if v)
    assert "—" not in text  # no em dashes in owner-facing text


def test_gave_up_venues_stay_on_waiting_tab_after_waiting_ones():
    rows = leadsheet.group_records([
        rec(1, venue_key="old", tier="A", lead_score=95, dba="FAKE OLD CLUB"),
        rec(2, venue_key="new", tier="A", lead_score=50, dba="FAKE NEW CLUB")])
    leadsheet.apply_contact(rows[0], check("gave_up", next_check_at=None), DAY)
    leadsheet.apply_contact(rows[1], check("waiting"), DAY)
    waiting = leadsheet.waiting_rows(rows, rows)
    assert [(r["business_name"], r["contact_status_text"]) for r in waiting] == [
        ("Fake New Club", "Waiting"), ("Fake Old Club", "Gave up")]


def test_contact_person_comes_from_the_filing_with_search_links_only():
    [row] = leadsheet.group_records([rec(
        1, source="wa_lcb_actions", state="WA", city="SEATTLE", dba="ZEBRA FAKE TAVERN",
        legal_name="ZEBRA FAKE LLC", raw={"Applicant(s)": "ZEBRA FAKE LLC; JANE Q TESTER"})])
    assert row["contact_person"] == "Jane Q Tester"
    assert row["person_linkedin_url"] == (
        "https://www.google.com/search?q=site%3Alinkedin.com%2Fin+%22Jane+Q+Tester%22"
        "+%22Seattle%22")
    assert row["person_instagram_url"] == (
        "https://www.google.com/search?q=site%3Ainstagram.com+%22Jane+Q+Tester%22")
    assert row["person_facebook_url"].endswith("%22Jane+Q+Tester%22+%22Seattle%22")
    [company_only] = leadsheet.group_records([rec(
        1, source="wa_lcb_actions", raw={"Applicant(s)": "QUOKKA HOLDINGS INC"})])
    assert company_only["contact_person"] is None
    assert company_only["person_linkedin_url"] is None
    [none] = leadsheet.group_records([rec(1)])  # TX: the filing names nobody
    assert none["contact_person"] is None
    wb = load_workbook(io.BytesIO(leadsheet.build_workbook([row])))
    headers, [cells] = _table(wb["New"])
    assert cells["Contact person"].value == "Jane Q Tester"
    assert cells["Person on LinkedIn"].value == "Search"


# --- Instagram found by web search (contact.py; invented handles) ---

def search_channel(handle, signals, score, label):
    return {"kind": "instagram_search", "value": "@" + handle,
            "url": f"https://www.instagram.com/{handle}/", "signals": signals,
            "score": score, "label": label}


def test_search_found_instagram_on_lead_tab_and_possible_on_waiting_tab():
    rows = leadsheet.group_records([
        rec(1, venue_key="found", tier="A", lead_score=80, dba="ZEBRA FAKE LOUNGE",
            event_type="new"),
        rec(2, venue_key="maybe", tier="A", lead_score=70, dba="QUOKKA FAKE CLUB",
            event_type="new")])
    by_key = {r["venue_key"]: r for r in rows}
    leadsheet.apply_contact(by_key["found"], check(
        "reachable", confidence_score=50, confidence_label="Likely",
        confidence_reason="Found by web search: name and city match",
        outreach_method="Instagram DM (found by search)", outreach_second=None,
        contact_kind="instagram_search", contact_value="@zebrafake.atx",
        contact_url="https://www.instagram.com/zebrafake.atx/",
        channels=[search_channel("zebrafake.atx", ["ig_search_name", "ig_search_city"],
                                 50, "Likely")]), DAY)
    leadsheet.apply_contact(by_key["maybe"], check("waiting", channels=[
        search_channel("quokkafake", ["ig_search_name", "ig_search_ambiguous"], 35,
                       "Unverified"),
        search_channel("quokkafake.atx", ["ig_search_name", "ig_search_ambiguous"], 35,
                       "Unverified"),
        {"kind": "filing_phone", "value": "(512) 555-0199", "url": "tel:+15125550199",
         "signals": ["filing_phone"], "score": 15, "label": "Unverified"}]), DAY)
    assert by_key["found"]["instagram_by_search"]
    assert not by_key["maybe"]["instagram_by_search"]
    assert by_key["maybe"]["verified_instagram_url"] is None

    wb = load_workbook(io.BytesIO(leadsheet.build_workbook(rows, rows)))
    _, cells = _table(wb["New"])
    [zebra] = cells
    assert zebra["Best way to reach"].value == "Instagram DM (found by search)"
    assert zebra["Contact"].value == "@zebrafake.atx"
    assert zebra["Contact"].hyperlink.target == "https://www.instagram.com/zebrafake.atx/"
    assert zebra["Why we trust it"].value == "Found by web search: name and city match"
    assert zebra["Instagram"].hyperlink.target == "https://www.instagram.com/zebrafake.atx/"

    w_headers, w_cells = _table(wb["Waiting on contact"], header_row=2)
    assert w_headers.index("Possible Instagram") == \
        w_headers.index("What we found (not verified)") + 1
    [quokka] = w_cells
    assert quokka["Possible Instagram"].value == "@quokkafake"
    assert quokka["Possible Instagram"].hyperlink.target == \
        "https://www.instagram.com/quokkafake/"
    assert quokka["Possible Instagram 2"].hyperlink.target == \
        "https://www.instagram.com/quokkafake.atx/"
    assert "Instagram (found by search): @quokkafake (Unverified 35)" in \
        quokka["What we found (not verified)"].value
    legend = [c.value for c in wb["How scoring works"]["A"]]
    assert "5. Instagram DM (found by search)" in legend
    assert any(str(v).startswith("Instagram found by web search") for v in legend)
    text = " ".join(str(v) for ws in wb for row in ws.values for v in row if v)
    assert "—" not in text


def test_possible_instagram_is_empty_without_search_finds():
    rows = checked_rows()
    wait = next(r for r in rows if r["venue_key"] == "wait")
    assert wait["possible_instagram_url"] is None and wait["possible_instagram_2_url"] is None
    wb = load_workbook(io.BytesIO(leadsheet.build_workbook(rows, rows)))
    [quokka] = _table(wb["Waiting on contact"], header_row=2)[1]
    assert quokka["Possible Instagram"].value is None


# --- owner names the filings publish (Source.contact "people") ---

def test_owner_names_from_five_more_sources():
    from licmon.sources.chicago_bacp import ChicagoBacpSource
    from licmon.sources.chicago_pending import ChicagoPendingSource
    from licmon.sources.ny_sla import NySlaSource
    from licmon.sources.tx_tabc import TxTabcSource

    assert TxTabcSource().contact({"owner": "JANE Q TESTER"}) == {"people": "JANE Q TESTER"}
    assert TxTabcSource().contact({"owner": "  "}) == {}
    assert CaAbcSource().contact({"rows": [{"Primary Name": "TESTER, JANE"},
                                           {"Primary Name": "TESTER, JANE"},
                                           {"Primary Name": "ZEBRA FAKE LLC"}]}) == {
        "people": "TESTER, JANE; ZEBRA FAKE LLC"}
    fl = FlAbtSource().contact({"Owner Name": "JOHN TESTER", "Mail Address 1": "9 TEST WAY",
                                "Mail City": "MIAMI", "Mail State": "FL", "Mail ZIP": "33101"})
    assert fl == {"people": "JOHN TESTER", "mailing_address": "9 TEST WAY, MIAMI, FL 33101"}
    assert NySlaSource().contact({"legalname": "QUOKKA FAKE INC"}) == {
        "people": "QUOKKA FAKE INC"}
    assert ChicagoBacpSource().contact({"legal_name": "JANE TESTER"}) == {
        "people": "JANE TESTER"}
    pending = ChicagoPendingSource()
    # The list only carries the popup link's label, never names: nothing.
    assert pending.contact({"Ownership": "Owners/Officers"}) == {}
    assert pending.contact({"Ownership": None}) == {}
    assert pending.contact({"Ownership": "JANE TESTER; JOHN TESTER"}) == {
        "people": "JANE TESTER; JOHN TESTER"}


def test_owner_who_is_a_person_fills_contact_person():
    [row] = leadsheet.group_records([rec(
        1, dba="ZEBRA FAKE LOUNGE", legal_name="JANE Q TESTER",
        raw={"owner": "JANE Q TESTER"})])
    assert row["people"] == "Jane Q Tester"
    assert row["contact_person"] == "Jane Q Tester"
    # A company owner, or a name that is just the business again, stays out.
    [row] = leadsheet.group_records([rec(
        2, dba="ZEBRA FAKE LOUNGE", legal_name="ZEBRA FAKE LLC",
        raw={"owner": "ZEBRA FAKE LLC"})])
    assert row["people"] is None and row["contact_person"] is None
    [row] = leadsheet.group_records([rec(
        3, source="ny_sla_pending", dba=None, legal_name="QUOKKA FAKE INC",
        raw={"legalname": "QUOKKA FAKE INC"})])
    assert row["contact_person"] is None


# --- Google's venue type moves the tier everywhere (contact.google_tier) ---

def test_name_only_a_is_a_club_or_lounge_name_without_other_signals():
    def only(**kw):
        [row] = leadsheet.group_records([rec(1, tier="A", **kw)])
        return row["name_only_a"]

    assert only(dba="ZEBRA FAKE LOUNGE") is True
    assert only(dba="ZEBRA FAKE COMEDY LOUNGE") is False  # ticketed-venue word
    assert only(dba="ZEBRA FAKE LOUNGE", license_type="LH") is False  # late hours license
    assert only(dba="ZEBRA FAKE LOUNGE", license_description="Cabaret") is False
    [b] = leadsheet.group_records([rec(1, tier="B", dba="ZEBRA FAKE LOUNGE")])
    assert b["name_only_a"] is False  # only A venues can be moved down


def test_tier_adjustment_shows_on_sheet_and_decides_attio():
    from licmon import attio

    by_key = {r["venue_key"]: r for r in leadsheet.group_records([
        rec(1, venue_key="down", tier="A", lead_score=55, hot=True, dba="ZEBRA FAKE LOUNGE"),
        rec(2, venue_key="up", tier="B", lead_score=45, dba="QUOKKA FAKE TAVERN"),
        rec(3, venue_key="same", tier="B", lead_score=62, dba="OKAPI FAKE PUB"),
    ])}
    down, up, same = by_key["down"], by_key["up"], by_key["same"]
    assert attio.eligible(down, 60) and not attio.eligible(up, 60)
    leadsheet.apply_contact(down, check("reachable", tier_adjusted="B",
                                        google_says="Google says restaurant",
                                        google_primary_type="mexican_restaurant"), DAY)
    leadsheet.apply_contact(up, check("reachable", tier_adjusted="A",
                                      google_says="Google says nightclub",
                                      google_primary_type="night_club"), DAY)
    leadsheet.apply_contact(same, check("reachable", google_primary_type="pub"), DAY)
    assert (down["priority"], down["priority_filed"], down["hot"]) == ("B", "A", "")
    assert down["business_type"] == "Bar / event venue"
    assert down["google_says"] == "Google says restaurant. Shown as B, filed as A"
    assert (up["priority"], up["priority_filed"]) == ("A", "B")
    assert up["google_says"] == "Google says nightclub. Shown as A, filed as B"
    assert (same["priority"], same["google_says"]) == ("B", "Pub")
    assert "priority_filed" not in same
    # Restaurants stop going to Attio as clubs; Google's nightclubs go.
    assert not attio.eligible(down, 60) and attio.eligible(up, 60)
    assert [r["venue_key"] for r in attio.candidates([down, up, same])] == ["same", "up"]
    assert attio.entry_values(up)["priority"] == "A"
    # Applying twice changes nothing more.
    leadsheet.apply_contact(down, check("reachable", tier_adjusted="B",
                                        google_says="Google says restaurant"), DAY)
    assert (down["priority"], down["priority_filed"]) == ("B", "A")


def test_platform_and_opening_columns_on_lead_tabs():
    rows = checked_rows()
    by_key = {r["venue_key"]: r for r in rows}
    leadsheet.apply_contact(by_key["reach"], check(
        "reachable", platforms=["SevenRooms", "Toast"], opening_soon=True,
        opening_signal="coming soon"), DAY)
    leadsheet.apply_contact(by_key["newly"], check(
        "reachable", newly_reachable_on=DAY, platforms=[], opening_signal="now open",
        opening_soon=False), DAY)
    wb = load_workbook(io.BytesIO(leadsheet.build_workbook(rows, rows)))
    headers, cells = _table(wb["New"])
    for header in ("Current platform", "Opening soon", "Google says", "Speakeasy / Attio"):
        assert header in headers
    got = {c["Business name"].value: (c["Current platform"].value, c["Opening soon"].value)
           for c in cells}
    assert got["Zebra Fake Lounge"] == ("SevenRooms, Toast", "Yes: coming soon")
    assert got["Fake Taproom"] == ("None found on their site", "Now open")
    speakeasy = leadsheet.apply_contact(dict(rows[0]), check(
        "reachable", platforms=["Speakeasy"]), DAY)
    assert speakeasy["on_speakeasy"] and speakeasy["pipeline"] == "Already on Speakeasy"


# --- Today's plan tab (outreach.build_plan) ---

def test_todays_plan_tab_comes_first_with_why_how_and_opener():
    rows = checked_rows()
    by_key = {r["venue_key"]: r for r in rows}
    by_key["reach"]["contact_person"] = "Jane Q Tester"
    leadsheet.apply_contact(by_key["reach"], check("reachable", platforms=[]), DAY)
    outcomes = [("contacted", "Call", 2), ("won", "Call", 1)]
    wb = load_workbook(io.BytesIO(leadsheet.build_workbook(rows, rows, outcomes=outcomes)))
    ws = wb.worksheets[0]
    assert ws.title == "Today's plan"
    assert ws["A1"].value.startswith("2 new venues to reach today, best first.")
    headers, cells = _table(ws, header_row=2)
    assert headers == ["Priority order", "Business name", "City", "Stage", "Why reach out",
                       "Best way to reach", "Contact", "Confidence", "Second best way",
                       "Opener", "Contact person", "Current platform", "Opening soon",
                       "Lead ID"]
    names = [c["Business name"].value for c in cells]
    assert names == ["Fake Taproom", "Zebra Fake Lounge"]  # newly reachable first
    zebra = cells[1]
    assert zebra["Priority order"].value == 2
    assert zebra["Why reach out"].value.splitlines()[0] == (
        "- License just filed, not open yet: pitch before launch")
    assert "- No ticketing or reservations found on their site: greenfield" in \
        zebra["Why reach out"].value
    assert zebra["Best way to reach"].value == "Instagram DM"
    assert zebra["Contact"].value == "@zebrafakelounge"
    assert zebra["Contact"].hyperlink.target == "https://www.instagram.com/zebrafakelounge/"
    assert zebra["Opener"].value.startswith("Hey Jane,\n\nCongrats on the upcoming opening")
    assert zebra["Opener"].alignment.wrap_text and zebra["Why reach out"].alignment.wrap_text
    assert zebra["Contact person"].value == "Jane Q Tester"
    assert zebra["Current platform"].value == "None found on their site"
    assert zebra["Lead ID"].value == "1"
    # The waiting venue and the venue never checked are not in the plan.
    assert "Quokka Fake Club" not in names and "Fake Bistro" not in names
    legend = {c.value: ws_b.value for c, ws_b in wb["How scoring works"].iter_rows(
        min_col=1, max_col=2)}
    assert legend["Outreach results, last 30 days"] == (
        "contacted 2, replied 0, won 1, wrong contact 0")
    assert legend["  Call"] == "contacted 2, won 1"
    text = " ".join(str(v) for row in wb["How scoring works"].values for v in row if v)
    assert "—" not in text


def test_todays_plan_tab_when_nobody_new_is_reachable():
    rows = [r for r in checked_rows() if r["venue_key"] in ("wait", "never")]
    wb = load_workbook(io.BytesIO(leadsheet.build_workbook(rows, rows)))
    ws = wb["Today's plan"]
    assert ws["A1"].value == "No new reachable venues today."
    assert ws.max_row == 2  # the line and the header
