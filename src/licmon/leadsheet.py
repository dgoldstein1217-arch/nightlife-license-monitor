"""The owner's lead spreadsheet: one clean row per venue, outreach first.

Built from the review queue, the contact details the official records
themselves publish (see Source.contact), and the contact lookup's results
(contact.py, table contact_checks) when it has run. People are never looked
up: the Contact person columns are the filing's own named person plus plain
search links the owner clicks by hand.

The workbook (build_workbook) has a New tab (new venues and unknown
history), an Existing venues tab (new owner or adding a permit), an All open
tab, one tab per state that has open leads, and a How scoring works tab.
Once the contact lookup has run, eligible venues with no verified contact
move from those tabs to a Waiting on contact tab, and the lead tabs gain
Best way to reach, Contact, Confidence and Why we trust it. Until it has
run, the workbook is laid out as before. Every tab sorts by lead score,
highest first (newly reachable venues first).

Local and email use only. Never print rows in GitHub Actions (public logs).
"""

from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime, timezone
from urllib.parse import quote_plus

from . import history as history_mod
from . import stage as stage_mod

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

#: (key, header, excel width). Order is the column order.
COLUMNS: list[tuple[str, str, int]] = [
    ("priority", "Priority", 9),
    ("hot", "Hot", 7),
    ("lead_score", "Score", 8),
    ("business_name", "Business name", 34),
    ("whats_new", "What's new", 16),
    ("company", "Company / owner", 34),
    ("business_type", "Business type", 32),
    ("filing", "Filing", 18),
    ("venue_history", "Venue history", 16),
    ("stage", "Stage", 11),
    ("filed_on", "Filed on", 12),
    ("phone", "Phone", 16),
    ("people", "Owner / applicant names", 34),
    ("contact_person", "Contact person", 24),
    ("person_linkedin_url", "Person on LinkedIn", 12),
    ("person_instagram_url", "Person on Instagram", 12),
    ("person_facebook_url", "Person on Facebook", 12),
    ("address", "Address", 34),
    ("city", "City", 18),
    ("state", "State", 7),
    ("zip", "ZIP", 8),
    ("market", "Market", 24),
    ("mailing_address", "Mailing address", 40),
    ("license", "License applied for", 40),
    ("map_url", "Map", 8),
    ("google_url", "Google", 9),
    ("instagram_url", "Instagram", 11),
    ("record_url", "Official record", 14),
    ("lead_ids", "Lead ID", 12),
]
HEADERS = [h for _, h, _ in COLUMNS]
#: All open and state tabs span many days: the queue date replaces What's new.
OPEN_COLUMNS = [("queue_date", "Queued on", 12) if key == "whats_new" else (key, h, w)
                for key, h, w in COLUMNS]

#: Lead tabs once the contact lookup has run: these go right after the name.
CONTACT_COLUMNS: list[tuple[str, str, int]] = [
    ("newly_reachable", "Newly reachable", 16),
    ("outreach_method", "Best way to reach", 28),
    ("contact_url", "Contact", 30),
    ("confidence", "Confidence", 14),
    ("confidence_reason", "Why we trust it", 54),
    ("outreach_second", "Second best way", 26),
]
FILING_PHONE_HEADER = "Filing phone (may be a lawyer)"
WAITING_TAB = "Waiting on contact"
NEWLY_REACHABLE = "Newly reachable"
#: The Waiting on contact tab: what was found, when it is checked next, and
#: the search links for a manual lookup. Possible Instagram links the
#: accounts the web search found that are not sure enough (name match only,
#: or two accounts that match equally) for the owner to open and confirm.
WAITING_COLUMNS: list[tuple[str, str, int]] = [
    ("priority", "Priority", 9),
    ("hot", "Hot", 7),
    ("lead_score", "Score", 8),
    ("business_name", "Business name", 34),
    ("contact_status_text", "Status", 10),
    ("found_text", "What we found (not verified)", 60),
    ("possible_instagram_url", "Possible Instagram", 18),
    ("possible_instagram_2_url", "Possible Instagram 2", 18),
    ("last_checked", "Last checked", 13),
    ("next_check", "Next check", 13),
    ("queue_date", "Queued on", 12),
    ("stage", "Stage", 11),
    ("company", "Company / owner", 34),
    ("contact_person", "Contact person", 24),
    ("person_linkedin_url", "Person on LinkedIn", 12),
    ("person_instagram_url", "Person on Instagram", 12),
    ("person_facebook_url", "Person on Facebook", 12),
    ("phone", FILING_PHONE_HEADER, 18),
    ("address", "Address", 34),
    ("city", "City", 18),
    ("state", "State", 7),
    ("zip", "ZIP", 8),
    ("market", "Market", 24),
    ("map_url", "Map", 8),
    ("google_url", "Google", 9),
    ("instagram_url", "Instagram", 11),
    ("record_url", "Official record", 14),
    ("lead_ids", "Lead ID", 12),
]


def with_contact_columns(columns: list[tuple[str, str, int]]) -> list[tuple[str, str, int]]:
    """A lead tab's columns once the contact lookup has run: CONTACT_COLUMNS
    after the name, a Facebook column after Instagram, and the filing phone
    labeled for what it is."""
    out = []
    for key, header, width in columns:
        if key == "phone":
            header = FILING_PHONE_HEADER
        out.append((key, header, width))
        if key == "business_name":
            out += CONTACT_COLUMNS
        if key == "instagram_url":
            out.append(("facebook_url", "Facebook", 11))
    return out


NEW_FILING = "New filing"
STAGE_ADVANCED = "Stage advanced"
DETAILS_CHANGED = "Details changed"

# Tier A/B mean nightclub or ticketed venue / bar-or-venue (see qualify.py);
# C is by license.
BUSINESS_TYPES = {
    "nightlife": "Restaurant / bar",
    "on_premise": "Restaurant",
    "hospitality_mfg": "Brewery / taproom / winery",
    "catering_event": "Caterer / event venue",
    "hotel": "Hotel",
}
_TYPE_ORDER = list(BUSINESS_TYPES)

# (pattern on the upper-cased application type, plain label). First match wins.
_FILINGS = [
    (r"NEW LICENSE", "Newly licensed"),
    (r"\bNEW\b|ORIGINAL", "New application"),
    (r"ASSUMPTION|TRANSFER|CHANGE OF (OWNER|CORP|OFFICER|STOCK)|OWNERSHIP", "Change of owner"),
    (r"LOCATION|RELOC|C_LOC", "New location"),
    (r"TRADENAME|TRADE NAME|NAME CHANGE", "Name change"),
    (r"CLASS|IN LIEU|UPGRADE|ADDED PRIVILEGE", "License upgrade"),
]


#: Added to Business type for adult venues (qualify.ADULT_WORDS). They stay in
#: the spreadsheet but never go to Attio or Slack.
ADULT_NOTE = " (adult)"


def _business_type(priority: str, cats: list, adult: bool = False) -> str:
    if priority == "A":
        kind = "Nightclub / lounge / ticketed venue"
    elif priority == "B":
        kind = ("Brewery / taproom" if cats and cats[0] == "hospitality_mfg"
                else "Bar / event venue")
    else:
        kind = BUSINESS_TYPES.get(cats[0], "Other") if cats else "Other"
    return kind + (ADULT_NOTE if adult else "")


def _search(base: str, query: str) -> str | None:
    """A plain search link. Instagram's own search needs a login, so the
    Instagram column is a Google search limited to instagram.com."""
    query = " ".join((query or "").split())
    return base + quote_plus(query) if query else None


def _filing(app_types: list[str], source: str) -> str:
    labels = []
    for t in app_types:
        up = t.upper()
        label = next((lab for pat, lab in _FILINGS if re.search(pat, up)), t.title())
        if label not in labels:
            labels.append(label)
    if not labels and source == "ca_abc_applications":
        return "New application"  # CA export lists pending applications only
    return ", ".join(labels)


def _stage(recs: list[dict], sources: dict) -> str:
    """Most advanced stage across a venue's records. Uses the stored stage,
    or the source's mapping for rows stored before stages existed."""
    found = []
    for rec in recs:
        value = rec.get("stage")
        if value is None:
            src = sources.get(rec.get("source"))
            value = (_source_stage(src, rec) if src
                     else stage_mod.from_status(rec.get("status")))
        found.append(value)
    return stage_mod.best(found) or ""


def _source_stage(src, rec: dict) -> str | None:
    from .models import Record

    try:
        return src.stage(Record(source=rec.get("source") or "", source_record_id="",
                                source_url="", status=rec.get("status"),
                                license_type=rec.get("license_type")))
    except Exception:  # noqa: BLE001 - a label only
        return None


def _whats_new(recs: list[dict]) -> str:
    """New filing > Stage advanced > Details changed, from the queue events."""
    types = {r.get("event_type") for r in recs}
    if types & {"new", "baseline"}:
        return NEW_FILING
    if any(r.get("event_type") == "changed" and "stage" in (r.get("changes") or {})
           for r in recs):
        return STAGE_ADVANCED
    return DETAILS_CHANGED if "changed" in types else ""


def _phone(value: str | None) -> str | None:
    digits = re.sub(r"\D", "", value or "")
    if len(digits) == 11 and digits.startswith("1"):
        digits = digits[1:]
    if len(digits) == 10:
        return f"({digits[:3]}) {digits[3:6]}-{digits[6:]}"
    return (value or "").strip() or None


def _title(text: str | None) -> str:
    """Tidy ALL-CAPS names; leave mixed case as the source wrote it."""
    text = " ".join((text or "").split())
    if text and text == text.upper() and any(c.isalpha() for c in text):
        text = text.title()
        text = re.sub(r"(\d)(St|Nd|Rd|Th)\b", lambda m: m.group(1) + m.group(2).lower(), text)
        text = re.sub(r"'S\b", "'s", text)
        text = re.sub(r"\b([A-Z][a-z])(?= \d{5})", lambda m: m.group(1).upper(), text)
    return re.sub(r"\b(Llc|Inc|Llp|Pllc|Ltd|Dba|Usa)\b", lambda m: m.group(1).upper(), text)


def _people(contacts: list[dict], names: tuple) -> str | None:
    """Owner/applicant names, minus entries that just repeat the business:
    the trade name (DBA), and the legal name when it is a company. A legal
    name that is a person (a sole proprietor) stays: Texas, New York,
    Chicago, California and Florida publish only the legal owner."""
    dba, legal = (tuple(names) + (None, None))[:2]
    skip = {re.sub(r"\W", "", n.upper()) for n in (dba, legal)
            if n and (n is dba and n != legal or _is_company(n))}
    people = []
    for c in contacts:
        for p in (c.get("people") or "").split(";"):
            p = _title(p.strip())
            if p and re.sub(r"\W", "", p.upper()) not in skip and p not in people:
                people.append(p)
    return "; ".join(people) or None


def _is_company(name: str) -> bool:
    words = set(re.sub(r"[^A-Z0-9]+", " ", name.upper()).split())
    return bool(words & history_mod._ENTITY_WORDS)


def _contact_person(people: str | None) -> str | None:
    """The first named person on the filing (not a company). Only what the
    filing says: people are never looked up."""
    for name in (people or "").split(";"):
        name = name.strip()
        if name and not _is_company(name):
            return name
    return None


def _person_links(person: str | None, city: str | None) -> dict:
    """Plain Google searches the owner opens by hand. Nothing is fetched."""
    if not person:
        return {"person_linkedin_url": None, "person_instagram_url": None,
                "person_facebook_url": None}
    where = f' "{city}"' if city else ""
    google = "https://www.google.com/search?q="
    return {
        "person_linkedin_url": _search(google, f'site:linkedin.com/in "{person}"{where}'),
        "person_instagram_url": _search(google, f'site:instagram.com "{person}"'),
        "person_facebook_url": _search(google, f'site:facebook.com "{person}"{where}'),
    }


def _uniq(values) -> list[str]:
    out: list[str] = []
    for v in values:
        v = (v or "").strip() if isinstance(v, str) else v
        if v and v not in out:
            out.append(v)
    return out


_SQL = """
SELECT q.queue_date, q.record_id, q.venue_key, q.tier, q.score, q.legal_name,
       q.dba, q.license_description, q.application_type, q.status,
       q.application_date, q.address, q.city, q.state, q.zip, q.metro,
       q.source, q.source_url, r.category, r.raw, q.event_type, q.changes,
       q.review_status, q.stage, q.lead_score, q.hot, q.license_type,
       q.first_seen_at, q.adult, q.venue_history, q.prior_licenses, q.prior_since
FROM review_queue q JOIN records r ON r.id = q.record_id
WHERE r.qualified
  AND (%(day)s::date IS NULL OR q.queue_date = %(day)s)
  AND (NOT %(open)s OR q.review_status = 'new')
  AND (%(keys)s::text[] IS NULL OR q.venue_key = ANY(%(keys)s))
ORDER BY q.queue_date DESC, q.score DESC, q.record_id
"""
_FIELDS = ["queue_date", "record_id", "venue_key", "tier", "score", "legal_name", "dba",
           "license_description", "application_type", "status", "application_date",
           "address", "city", "state", "zip", "metro", "source", "source_url",
           "category", "raw", "event_type", "changes", "review_status", "stage",
           "lead_score", "hot", "license_type", "first_seen_at", "adult",
           "venue_history", "prior_licenses", "prior_since"]


def load_rows(conn, day: date | None, open_only: bool = False,
              venue_keys: list[str] | None = None, with_contacts: bool = True) -> list[dict]:
    """Queued leads for `day` (all days if None), one row per venue per day.

    `venue_keys` limits it to those venues, latest queue day each. With
    `with_contacts`, each row gets its contact lookup result, and a day's
    rows also include venues that became reachable that day although they
    were queued earlier (open ones only)."""
    with conn.cursor() as cur:
        cur.execute(_SQL, {"day": day, "open": open_only, "keys": venue_keys})
        records = [dict(zip(_FIELDS, r)) for r in cur.fetchall()]
    rows = group_records(records)
    if venue_keys is not None:
        latest: dict[str, dict] = {}
        for row in rows:
            key = row["venue_key"]
            if key not in latest or row["queue_date"] > latest[key]["queue_date"]:
                latest[key] = row
        rows = sorted(latest.values(), key=sort_key)
    if not with_contacts or not _has_contact_table(conn):
        return rows
    if day is not None:
        have = {r["venue_key"] for r in rows}
        with conn.cursor() as cur:
            cur.execute("SELECT venue_key FROM contact_checks WHERE newly_reachable_on = %s "
                        "AND status = 'reachable'", (day,))
            extra = [k for (k,) in cur.fetchall() if k not in have]
        if extra:
            rows += load_rows(conn, None, open_only=True, venue_keys=extra,
                              with_contacts=False)
    attach_contacts(conn, rows, day)
    rows.sort(key=sort_key)
    return rows


def _has_contact_table(conn) -> bool:
    with conn.cursor() as cur:
        cur.execute("SELECT to_regclass('contact_checks') IS NOT NULL")
        return bool(cur.fetchone()[0])


_CHECK_FIELDS = ["venue_key", "status", "channels", "confidence_score", "confidence_label",
                 "confidence_reason", "outreach_method", "outreach_second", "contact_kind",
                 "contact_value", "contact_url", "maps_url", "last_checked_at",
                 "next_check_at", "newly_reachable_on"]


def attach_contacts(conn, rows: list[dict], day: date | None = None) -> None:
    """Add each row's contact lookup result (apply_contact). Rows the lookup
    never checked are left as they are."""
    keys = sorted({r["venue_key"] for r in rows if r.get("venue_key")})
    if not keys:
        return
    with conn.cursor() as cur:
        cur.execute(f"SELECT {', '.join(_CHECK_FIELDS)} FROM contact_checks "
                    "WHERE venue_key = ANY(%s)", (keys,))
        checks = {r[0]: dict(zip(_CHECK_FIELDS, r)) for r in cur.fetchall()}
    for row in rows:
        check = checks.get(row.get("venue_key"))
        if check:
            apply_contact(row, check, day)


def apply_contact(row: dict, check: dict, day: date | None = None) -> dict:
    """Put one contact_checks row on a lead row: the status, best way to
    reach, the contact as a link, confidence, the reason, verified links for
    the lead tabs and a summary of unverified finds for the Waiting tab.
    `day` is the sheet's day (today UTC if None) for the Newly reachable flag."""
    from . import contact

    day = day or datetime.now(timezone.utc).date()
    status = check.get("status")
    channels = []
    for c in check.get("channels") or []:
        known = {k: c.get(k) for k in ("kind", "value", "url", "signals", "score", "label",
                                         "last_post")}
        channels.append(contact.Channel(**{k: v for k, v in known.items() if v is not None}))
    best = contact.best_channels([c for c in channels if contact.reachable(c)])
    found = sorted((c for c in channels if not contact.reachable(c)), key=lambda c: -c.score)
    is_reachable = status == contact.REACHABLE

    def url(kind):
        ch = best.get(kind)
        return ch.url if ch else None

    possible = sorted((c for c in channels
                       if c.kind == contact.INSTAGRAM_SEARCH and not contact.reachable(c)),
                      key=lambda c: -c.score)[:2]
    possible += [None] * (2 - len(possible))

    row.update({
        "contact_status": status,
        "contact_status_text": contact.STATUS_TEXT.get(status, status or ""),
        "outreach_method": check.get("outreach_method"),
        "outreach_second": check.get("outreach_second"),
        "contact_url": check.get("contact_url") if is_reachable else None,
        "contact_url_text": check.get("contact_value") if is_reachable else None,
        "confidence": f"{check.get('confidence_label')} {check.get('confidence_score')}",
        "confidence_reason": check.get("confidence_reason"),
        "newly_reachable": (NEWLY_REACHABLE if is_reachable
                            and check.get("newly_reachable_on") == day else ""),
        "last_checked": check.get("last_checked_at"),
        "next_check": check.get("next_check_at"),
        "found_text": "; ".join(f"{contact.KIND_TEXT.get(c.kind, c.kind)}: {c.value} "
                                f"({c.label} {c.score})" for c in found)
                      or "Nothing found yet",
        "verified_phone": best[contact.PHONE].value if contact.PHONE in best else None,
        "verified_email": best[contact.EMAIL].value if contact.EMAIL in best else None,
        "verified_website_url": url(contact.WEBSITE),
        "verified_instagram_url": url(contact.INSTAGRAM) or url(contact.INSTAGRAM_SEARCH),
        # The verified Instagram came from the web search, not the website.
        "instagram_by_search": (contact.INSTAGRAM not in best
                                and contact.INSTAGRAM_SEARCH in best),
        "possible_instagram_url": possible[0].url if possible[0] else None,
        "possible_instagram_url_text": possible[0].value if possible[0] else None,
        "possible_instagram_2_url": possible[1].url if possible[1] else None,
        "possible_instagram_2_url_text": possible[1].value if possible[1] else None,
        "verified_facebook_url": url(contact.FACEBOOK),
    })
    return row


def group_records(records: list[dict]) -> list[dict]:
    """Merge application records (best score first) into venue rows."""
    from .sources import all_sources

    by_source = {s.name: s for s in all_sources()}
    groups: dict[tuple, list[dict]] = {}
    for rec in records:
        groups.setdefault((rec["queue_date"], rec["venue_key"]), []).append(rec)

    rows = []
    for (queue_date, _), recs in groups.items():
        recs.sort(key=lambda r: (-(r.get("score") or 0), r["record_id"]))
        top = recs[0]
        contacts = []
        for rec in recs:
            src = by_source.get(rec["source"])
            try:
                contacts.append(src.contact(rec.get("raw") or {}) if src else {})
            except Exception:  # noqa: BLE001 - contact details are best effort
                contacts.append({})
        dates = [r["application_date"] for r in recs if r.get("application_date")]
        cats = sorted({r.get("category") for r in recs},
                      key=lambda c: _TYPE_ORDER.index(c) if c in _TYPE_ORDER else 99)
        name = top.get("dba") or top.get("legal_name")
        company = top.get("legal_name") if top.get("legal_name") != name else None
        address = ", ".join(p for p in (top.get("address"), top.get("city"),
                                        top.get("state"), top.get("zip")) if p)
        map_q = " ".join(p for p in (name, address) if p)
        web_q = " ".join(p for p in (name, top.get("city"), top.get("state")) if p)
        seen = [r["first_seen_at"] for r in recs if r.get("first_seen_at")]
        since = [r["prior_since"] for r in recs if r.get("prior_since")]
        people = _people(contacts, (top.get("dba") or None, top.get("legal_name")))
        person = _contact_person(people)
        rows.append({
            "queue_date": queue_date,
            "venue_key": top.get("venue_key"),
            "priority": min((r.get("tier") or "C") for r in recs),
            "score": top.get("score") or 0,
            "lead_score": max((r.get("lead_score") or 0) for r in recs),
            "hot": "Hot" if any(r.get("hot") for r in recs) else "",
            "adult": any(r.get("adult") for r in recs),
            "whats_new": _whats_new(recs),
            "first_seen": min(seen) if seen else None,
            "business_name": _title(name),
            "company": _title(company) or None,
            "business_type": _business_type(min((r.get("tier") or "C") for r in recs), cats,
                                            any(r.get("adult") for r in recs)),
            "filing": _filing(_uniq(r.get("application_type") for r in recs), top["source"]),
            # Not checked yet (stored before the check existed) reads Unknown.
            "venue_history": history_mod.best(r.get("venue_history") or history_mod.UNKNOWN
                                              for r in recs),
            "prior_licenses": max((r.get("prior_licenses") or 0) for r in recs),
            "prior_since": min(since) if since else None,
            "stage": _stage(recs, by_source),
            "filed_on": min(dates) if dates else None,
            "phone": ", ".join(_uniq(_phone(c.get("phone")) for c in contacts)) or None,
            "people": people,
            "contact_person": person,
            **_person_links(person, _title(top.get("city"))),
            "address": _title(top.get("address")),
            "city": _title(top.get("city")),
            "state": (top.get("state") or "").upper(),
            "zip": top.get("zip") or "",
            "market": top.get("metro") or "",
            "mailing_address": _title(next((c["mailing_address"] for c in contacts
                                            if c.get("mailing_address")), None)) or None,
            "license": "; ".join(_uniq(r.get("license_description") for r in recs)),
            "map_url": _search("https://www.google.com/maps/search/?api=1&query=", map_q),
            "google_url": _search("https://www.google.com/search?q=", web_q),
            "instagram_url": _search("https://www.google.com/search?q=",
                                     f"site:instagram.com {name}" if name else ""),
            "record_url": top.get("source_url"),
            "lead_ids": " ".join(str(r["record_id"]) for r in
                                 sorted(recs, key=lambda r: r["record_id"])),
        })
    rows.sort(key=sort_key)
    return rows


def sort_key(row: dict):
    """Newly reachable first, then highest lead score, priority, market, name."""
    return (not row.get("newly_reachable"), -(row.get("lead_score") or 0),
            row.get("priority") or "C",
            row.get("market") or "", row.get("business_name") or "")


def _cell(value):
    if isinstance(value, datetime):
        return value.date()
    return "" if value is None else value


def write_csv(rows: list[dict], out) -> None:
    w = csv.writer(out)
    w.writerow(HEADERS)
    for row in rows:
        w.writerow([_csv_safe(_cell(row.get(k))) for k, _, _ in COLUMNS])


def _csv_safe(value):
    """Stop spreadsheet apps from running a source string as a formula."""
    if isinstance(value, str) and value[:1] in ("=", "+", "@"):
        return "'" + value
    return value


def build_xlsx(rows: list[dict], title: str = "Leads") -> bytes:
    """One sheet of venue rows (the New tab layout)."""
    from openpyxl import Workbook

    wb = Workbook()
    _write_sheet(wb.active, rows, COLUMNS, title)
    return _save(wb)


EXISTING_NOTE = ("These venues have been open before. They are changing owner or adding "
                 "a permit, so they are not new venues.")


def is_existing(row: dict) -> bool:
    """New owner or Adding a permit: the venue has been around."""
    return row.get("venue_history") in history_mod.EXISTING


WAITING_NOTE = ("No verified contact yet. We check each venue again every week. Once we "
                "find a way to reach it, it moves back to the lead tabs as Newly reachable. "
                "After 120 days it says Gave up. Use the search links to look one up "
                "yourself.")


def contact_checked(rows: list[dict]) -> bool:
    """Has the contact lookup looked at any of these venues?"""
    return any(r.get("contact_status") for r in rows)


def is_waiting(row: dict) -> bool:
    """Checked, and no verified contact (waiting or gave up)."""
    return row.get("contact_status") in ("waiting", "gave_up")


def lead_tab_view(row: dict) -> dict:
    """A checked venue on a lead tab: verified links in place of the search
    links (those stay on the Waiting tab, where they help a manual lookup).
    Venues the lookup never checked keep their search links."""
    if not row.get("contact_status"):
        return row
    out = dict(row)
    out["google_url"] = row.get("verified_website_url")
    out["google_url_text"] = "Website"
    out["instagram_url"] = row.get("verified_instagram_url")
    out["instagram_url_text"] = "Profile"
    out["facebook_url"] = row.get("verified_facebook_url")
    out["facebook_url_text"] = "Page"
    return out


def waiting_rows(*groups: list[dict]) -> list[dict]:
    """Waiting and gave-up venues from all the given rows, once each (latest
    queue day), waiting before gave up, then by score."""
    latest: dict[str, dict] = {}
    for rows in groups:
        for row in rows:
            if not is_waiting(row):
                continue
            key = row.get("venue_key")
            have = latest.get(key)
            if have is None or (row.get("queue_date") or date.min) > (have.get("queue_date")
                                                                        or date.min):
                latest[key] = row
    return sorted(latest.values(),
                  key=lambda r: (r.get("contact_status") == "gave_up",) + sort_key(r))


def build_workbook(new_rows: list[dict], open_rows: list[dict] | None = None) -> bytes:
    """The owner's workbook: New (new venues and unknown history), Existing
    venues (new owner or adding a permit), All open, one tab per state with
    open leads (from the data, not a fixed list), then How scoring works.

    Once the contact lookup has run (any row checked), venues it checked
    without finding verified contact leave those tabs for a Waiting on
    contact tab after Existing venues, and the lead tabs gain the contact
    columns. Venues it never checked stay where they were."""
    from openpyxl import Workbook

    open_rows = sorted(open_rows or [], key=sort_key)
    new_rows = sorted(new_rows, key=sort_key)
    checked = contact_checked(new_rows) or contact_checked(open_rows)
    columns, open_columns = COLUMNS, OPEN_COLUMNS
    waiting: list[dict] = []
    if checked:
        columns, open_columns = with_contact_columns(COLUMNS), with_contact_columns(OPEN_COLUMNS)
        waiting = waiting_rows(new_rows, open_rows)
        new_rows = [lead_tab_view(r) for r in new_rows if not is_waiting(r)]
        open_rows = [lead_tab_view(r) for r in open_rows if not is_waiting(r)]
    wb = Workbook()
    _write_sheet(wb.active, [r for r in new_rows if not is_existing(r)], columns, "New")
    _write_sheet(wb.create_sheet(), [r for r in new_rows if is_existing(r)], columns,
                 "Existing venues", note=EXISTING_NOTE)
    if checked:
        _write_sheet(wb.create_sheet(), waiting, WAITING_COLUMNS, WAITING_TAB,
                     note=WAITING_NOTE)
    _write_sheet(wb.create_sheet(), open_rows, open_columns, "All open")
    for state in sorted({r.get("state") or "Other" for r in open_rows}):
        _write_sheet(wb.create_sheet(),
                     [r for r in open_rows if (r.get("state") or "Other") == state],
                     open_columns, state)
    _write_legend(wb.create_sheet(), contact=checked)
    return _save(wb)


def _save(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _write_sheet(ws, rows: list[dict], columns: list[tuple[str, str, int]],
                 title: str, note: str | None = None) -> None:
    """One tab. A `note` goes in row 1 above the header."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ws.title = re.sub(r"[\[\]:*?/\\]", "-", title)[:31] or "Leads"
    top = 1
    if note:
        ws.append([note])
        ws.cell(row=1, column=1).font = Font(italic=True)
        top = 2
    ws.append([h for _, h, _ in columns])
    head_fill = PatternFill("solid", fgColor="1F2937")
    for cell in ws[top]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[top].height = 22
    fills = {"A": PatternFill("solid", fgColor="D1FAE5"),
             "B": PatternFill("solid", fgColor="FEF3C7"),
             "C": PatternFill("solid", fgColor="F3F4F6")}
    hot_fill = PatternFill("solid", fgColor="FEE2E2")
    link_font = Font(color="1D4ED8", underline="single")
    # A row may carry "<key>_text" to show instead of the default word.
    link_cols = {"map_url": "Map", "google_url": "Search", "instagram_url": "Search",
                 "record_url": "Record", "facebook_url": "Page", "contact_url": "Open",
                 "person_linkedin_url": "Search", "person_instagram_url": "Search",
                 "person_facebook_url": "Search", "possible_instagram_url": "Profile",
                 "possible_instagram_2_url": "Profile"}
    keys = [key for key, _, _ in columns]
    for i, row in enumerate(rows, start=top + 1):
        for j, key in enumerate(keys, start=1):
            value = _cell(row.get(key))
            cell = ws.cell(row=i, column=j)
            if key in link_cols:
                if isinstance(value, str) and value.startswith(("http://", "https://",
                                                                "mailto:", "tel:")):
                    cell.value = str(row.get(key + "_text") or link_cols[key])
                    if cell.data_type == "f":
                        cell.data_type = "s"  # a website's text is never a formula
                    cell.hyperlink = value
                    cell.font = link_font
                continue
            cell.value = value
            if cell.data_type == "f":
                cell.data_type = "s"  # never let a source string become a formula
            if isinstance(value, date):
                cell.number_format = "yyyy-mm-dd"
        prio = ws.cell(row=i, column=keys.index("priority") + 1)
        prio.alignment = Alignment(horizontal="center")
        if row.get("priority") in fills:
            prio.fill = fills[row["priority"]]
            prio.font = Font(bold=True)
        if row.get("hot"):
            hot = ws.cell(row=i, column=keys.index("hot") + 1)
            hot.fill = hot_fill
            hot.font = Font(bold=True, color="B91C1C")
    for j, (_, _, width) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(j)].width = width
    ws.freeze_panes = f"{get_column_letter(keys.index('business_name') + 2)}{top + 1}"
    ws.auto_filter.ref = (f"A{top}:{get_column_letter(len(columns))}"
                          f"{max(top, len(rows) + top)}")


def contact_legend_lines() -> list[tuple[str, str]]:
    """The contact part of How scoring works, from contact.py's own tables."""
    from . import contact as c

    pts = c.POINTS
    lines = [
        ("", ""),
        ("Contact confidence", "Points"),
        ("Google listing at the same address as the filing", pts["listing_address"]),
        ("The listing's name matches the venue name", f"{pts['listing_name']}. Without it, "
         "anything from that listing stays Unverified (it may be the old business)."),
        ("Phone: Google says the place is open", pts["listing_open"]),
        ("Phone: the venue website shows the same number", pts["phone_on_website"]),
        ("Website: the site loads", pts["website_loads"]),
        ("Email or Instagram linked from the venue website", pts["linked_from_website"]),
        ("Facebook page linked from the venue website",
         f"{pts['facebook_linked']}. Facebook cannot be checked further, so at most Likely."),
        ("Instagram links back to the same website", pts["ig_links_back"]),
        ("Instagram bio names the street address or ZIP", pts["ig_bio_address"]),
        ("Instagram bio names the city or neighborhood", pts["ig_bio_city"]),
        ("Instagram the API cannot read (personal account)", pts["ig_unreadable"]),
        ("Instagram found by web search: its name or handle matches the venue",
         f"{pts['ig_search_name']}. Only when the website links no Verified or Likely "
         "Instagram. Without the name, a result never counts."),
        ("Found by web search: the snippet names the street address or ZIP",
         pts["ig_search_address"]),
        ("Found by web search: the city or metro (like ATX or NYC) in the result",
         f"{pts['ig_search_city']}. Name and address, or name and city, is Likely at "
         f"most ({c.CAP_SEARCH}): a venue with the same name in the same city would "
         "look the same, so check the profile first."),
        ("Found by web search: name only, or two accounts match equally",
         "Unverified. Shown in Possible Instagram on the Waiting on contact tab to "
         "open and confirm by hand."),
        ("Phone from the license filing", f"{pts['filing_phone']}. Often a lawyer or "
         "expediter, so it never makes a venue reachable on its own."),
        ("Confidence labels", f"Verified {c.LABEL_MIN[c.VERIFIED]} or more, Likely "
                   f"{c.LABEL_MIN[c.LIKELY]} to {c.LABEL_MIN[c.VERIFIED] - 1}, Unverified "
                   f"below {c.LABEL_MIN[c.LIKELY]}, None when nothing was found. A venue "
                   "is reachable when one way to reach it is Verified or Likely."),
        ("Google says it closed for good", "0 for everything from that listing."),
        ("", ""),
        ("Best way to reach", "First that fits"),
    ]
    lines += [(f"{n}. {method}", text) for n, (method, text) in enumerate(c.METHOD_TEXT, 1)]
    lines += [
        (c.WAIT_METHOD, "Nothing Verified or Likely yet. The venue is on the Waiting on "
                        "contact tab and is checked again every week, for 120 days."),
        (NEWLY_REACHABLE, "Was waiting, and a recheck found a verified way to reach it. "
                          "Shown first."),
        ("Contact person", "The person named on the filing, if any. The search links are "
                           "plain searches; nobody is looked up automatically, and a "
                           "person never counts toward contact confidence."),
    ]
    return lines


def legend_lines(contact: bool = False) -> list[tuple[str, str]]:
    """(label, value) rows for the How scoring works tab, built from the
    live points table in qualify.py so the two never drift. `contact` adds
    the contact confidence and outreach rules (once the lookup has run)."""
    from . import qualify as q

    lines = [
        ("How scoring works", ""),
        ("", ""),
        ("Score", "0 to 100. Higher means a better ticketing fit. Parts add up."),
        ("Hot", f"A priority A venue with a score of {q.hot_min_score()} or more. "
                "Adult venues are never Hot."),
        ("", ""),
        ("Venue type", "Points"),
        ("A: nightclubs, lounges and ticketed venues (comedy, live music, sports, "
         "theaters, event venues)", q.TIER_POINTS["A"]),
        ("B: bars, karaoke, billiards, bowling, cinemas; a lounge on a restaurant",
         q.TIER_POINTS["B"]),
        ("C: restaurant", q.TIER_POINTS["C"]),
        ("", ""),
        ("Nightlife license (highest one counts)", "Points"),
    ]
    lines += [(label, pts) for pts, label in sorted(q.NIGHTLIFE_LICENSE_POINTS.values(),
                                                   key=lambda v: -v[0])]
    lines += [("", ""), ("Stage", "Points")]
    lines += [(s, q.STAGE_POINTS[s]) for s in stage_mod.STAGES]
    lines += [
        ("", ""),
        ("Filing type", "Points"),
        ("New filing or new location", q.FILING_POINTS[0][1]),
        ("Change of owner", next(p for _, p, lab in q.FILING_POINTS
                                 if lab == "change of owner")),
        ("", ""),
        ("Venue history", "Points"),
        (history_mod.NEW_VENUE, "0. No license at this address before, or the last one "
                                "ended more than two years ago."),
        (history_mod.NEW_OWNER, f"{history_mod.SCORE_ADJUST[history_mod.NEW_OWNER]}. A "
                                "different company holds or recently held a license here. "
                                "Never Hot. Still goes to Attio."),
        (history_mod.ADDING_PERMIT, f"{history_mod.SCORE_ADJUST[history_mod.ADDING_PERMIT]}. "
                                    "The same company already holds a license here. Never "
                                    "Hot, never sent to Attio or named in Slack."),
        (history_mod.UNKNOWN, "0. The state publishes no license list we can check "
                              "(Washington, except its own ownership and class filings), "
                              "or the address has no street number or ZIP."),
        ("", ""),
        ("Stages", ""),
        ("Licensed", "Issued or active. Florida counts only if issued in the last 60 days."),
        ("Approved", "Approved or conditional, not yet active."),
        ("In review", "Past intake, in process."),
        ("Received", "Just filed."),
        ("", ""),
        ("Labels", ""),
        ("(adult) in Business type", "Gentlemen's club, strip club or similar. Kept here, "
                                     "never sent to Attio or named in Slack."),
        ("Chicago amusement license", "Adds points. A bar name with one stays B."),
        ("Stadium concessionaire", "A food and drink operator (Levy, Aramark and others) "
                                   "at a named stadium, arena or venue counts as A."),
        ("", ""),
        ("What's new", ""),
        (NEW_FILING, "First time this filing showed up."),
        (STAGE_ADVANCED, "A filing we already had moved to a later stage."),
        (DETAILS_CHANGED, "Name, address, license or status changed."),
        ("", ""),
        ("Tabs", ""),
        ("New", "The day's leads: new venues and unknown history."),
        ("Existing venues", "The day's leads at venues that have been open before: "
                            "new owner or adding a permit."),
        ("All open", "Every lead not yet reviewed, from all days."),
        ("State tabs", "All open, split by state."),
    ]
    if contact:
        lines.insert(len(lines) - 1, (WAITING_TAB, "Venues we cannot reach yet: what was "
                                                   "found, when they are checked again, and "
                                                   "search links."))
        lines += contact_legend_lines()
    return lines


def _write_legend(ws, contact: bool = False) -> None:
    from openpyxl.styles import Font

    ws.title = "How scoring works"
    for label, value in legend_lines(contact):
        ws.append([label, value])
        if value in ("Points", "First that fits") or label in (
                "How scoring works", "Stages", "Labels", "What's new", "Tabs"):
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
            ws.cell(row=ws.max_row, column=2).font = Font(bold=True)
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 80
